from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
import json
import os

from openpyxl import Workbook

from artifact_tools import ValidationFailure, validate_results_file, write_bytes_atomic
from xlsx_support import add_table, workbook_bytes


STATUS_LABELS = {"passed": "通过", "failed": "失败", "blocked": "阻塞", "flaky": "不稳定", "not-run": "未执行", "not-evaluated": "未检查"}


def automatic_gate(results, manifest):
    cases = {case["id"]: case for case in manifest["cases"]}
    relevant = [r for r in results["results"] if next(a for a in cases[r["case_id"]]["automation"] if a["target_id"] == r["target_id"])["feasibility"] != "manual"]
    if not relevant:
        return "NO_AUTOMATABLE_CASES"
    if any(r["status"] == "failed" for r in relevant):
        return "FAILED"
    if any(r["status"] != "passed" for r in relevant):
        return "INCOMPLETE"
    if any(not r.get("native_test") for r in relevant):
        return "NO_CODE_EXECUTION"
    return "PASSED"


def render_report(results, manifest, *, results_dir=None, report_dir=None) -> bytes:
    run = results["run"]
    cases = {case["id"]: case for case in manifest["cases"]}
    workbook = Workbook()
    workbook.remove(workbook.active)
    counts = Counter(r["status"] for r in results["results"])
    assertion_rows, result_rows, evidence_rows = [], [], []
    planned_auto = evaluated_auto = passed_auto = 0
    def report_link(path):
        if results_dir is None or report_dir is None:
            return str(path)
        resolved = (results_dir / path).resolve()
        try:
            return Path(os.path.relpath(resolved, report_dir)).as_posix()
        except ValueError:
            return str(resolved)
    for result in results["results"]:
        case = cases[result["case_id"]]
        target = result["target_id"]
        assessment = next(a for a in case["automation"] if a["target_id"] == target)
        actuals = {a["assertion_id"]: a for a in result["assertions"]}
        eligible = assessment["feasibility"] != "manual"
        for planned in case["assertions"]:
            if target not in planned["target_ids"]:
                continue
            actual = actuals.get(planned["id"], {})
            status = actual.get("status", "not-evaluated")
            if eligible:
                planned_auto += 1
                evaluated_auto += status != "not-evaluated"
                passed_auto += status == "passed"
            assertion_rows.append([case["id"], target, planned["id"], planned["description"], planned["observation"], planned["expected"], actual.get("actual", ""), STATUS_LABELS[status], "\n".join(actual.get("evidence_refs", [])), actual.get("reason", result.get("blocker", ""))])
            assertion_rows[-1].extend([planned["check"]["kind"], planned["check"]["timing"], json.dumps(actual.get("observations", []), ensure_ascii=False), json.dumps(actual.get("judgments", []), ensure_ascii=False)])
        result_rows.append([case["id"], case["title"], target, STATUS_LABELS[result["status"]], assessment["feasibility"], result["attempts"], result["tool"], result.get("observed", ""), result.get("failure_reason", result.get("blocker", ""))])
        for evidence in result["evidence"]:
            evidence_rows.append([case["id"], target, evidence["id"], evidence["type"], evidence["description"], report_link(evidence["path"]), evidence["sha256"]])
    summary = [["项目", manifest["project"]["name"]], ["执行门禁", automatic_gate(results, manifest)], ["运行编号", run["run_id"]], ["用例集", run["suite"]], ["目标", "\n".join(run["targets"])], ["环境", run["environment"]["name"]], ["构建", run["environment"]["build"]], ["开始", run["started_at"]], ["结束", run["finished_at"]], ["用例基线 SHA256", run["manifest_sha256"]], ["用例 Excel SHA256", run["case_workbook_sha256"]], ["执行条件 SHA256", run["execution_profile_sha256"]]]
    summary += [[label, counts[status]] for status, label in STATUS_LABELS.items() if status != "not-evaluated"]
    summary += [["自动化必检断言数", planned_auto], ["已检查断言数", evaluated_auto], ["通过断言数", passed_auto], ["断言检查覆盖率", f"{evaluated_auto}/{planned_auto}" if planned_auto else "不适用"], ["说明", "人工用例和阻塞项保留在明细；门禁仅反映已选自动化范围，不能证明整个项目无缺陷。"]]
    if run.get("notes"):
        summary.append(["执行说明", run["notes"]])
    add_table(workbook, "执行汇总", ["项目", "内容"], summary, {"项目": 30, "内容": 100})
    add_table(workbook, "用例结果", ["用例编号", "用例名称", "目标编号", "状态", "自动化评估", "尝试次数", "执行工具", "实际观察", "失败或阻塞原因"], result_rows, {"用例名称": 40, "实际观察": 60, "失败或阻塞原因": 60})
    implementations = []
    for result in results["results"]:
        if not result.get("native_test"):
            continue
        native = result["native_test"]
        code, native_report = native["asset"], ""
        if results_dir is not None:
            from artifact_tools import load_json
            receipt = load_json(results_dir / native["receipt_path"])
            asset = next((a for a in receipt["assets"] if a["path"].replace("\\", "/") == native["asset"].replace("\\", "/")), None)
            if asset is None:
                raise ValidationFailure(["native test has no archived executable asset: " + native["asset"]])
            code = report_link(asset["evidence_path"])
            report = receipt.get("human_report", receipt.get("report"))
            if report:
                native_report = report_link(report["path"])
        implementations.append([result["case_id"],result["target_id"],result["tool"],native["route"],code,native["test_id"],report_link(native["receipt_path"]),native_report])
    implementation_sheet = add_table(workbook,"自动化实现",["用例编号","目标编号","运行框架","证据路线","测试代码快照","框架测试编号","执行记录","框架原生报告"],implementations,{"测试代码快照":60,"框架测试编号":65,"执行记录":65,"框架原生报告":65})
    for row in range(2, implementation_sheet.max_row + 1):
        for column in (5,7,8):
            cell = implementation_sheet.cell(row,column)
            if cell.value:
                cell.hyperlink = str(cell.value)
                cell.style = "Hyperlink"
    add_table(workbook, "断言结果", ["用例编号", "目标编号", "断言编号", "检查内容", "观察位置", "预期结果", "实际结果", "状态", "证据编号", "未检查原因", "比较方式", "检查时机", "各次观察引用", "证据评判理由"], assertion_rows, {"预期结果": 50, "实际结果": 50, "未检查原因": 50})
    sheet = add_table(workbook, "证据索引", ["用例编号", "目标编号", "证据编号", "类型", "说明", "文件路径", "SHA256"], evidence_rows, {"说明": 55, "文件路径": 65})
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, 6)
        cell.hyperlink = str(cell.value)
        cell.style = "Hyperlink"
    return workbook_bytes(workbook)


def main():
    parser = argparse.ArgumentParser(description="Render a validated Wewo QA execution workbook.")
    parser.add_argument("results", type=Path)
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    try:
        results_path, manifest_path = args.results.resolve(), args.manifest.resolve()
        output = args.output.resolve() if args.output else results_path.parent / "test-execution-report.xlsx"
        from artifact_layout import SUITES, case_workbook_path
        if output in {results_path, manifest_path, *(case_workbook_path(manifest_path, s) for s in SUITES)}:
            raise ValidationFailure(["report must not overwrite an execution or case baseline"])
        results, manifest = validate_results_file(results_path, manifest_path)
        write_bytes_atomic(output, render_report(results, manifest, results_dir=results_path.parent, report_dir=output.parent))
    except (ValidationFailure, OSError) as exc:
        print(f"ERROR: {exc}")
        return 1
    print(f"WROTE: {output}")
    return 0

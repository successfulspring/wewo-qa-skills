from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from artifact_tools import (
    ValidationFailure, sha256_file, validate_manifest_data, validate_manifest_file,
    write_bytes_atomic, write_text_atomic,
)
from artifact_layout import SUITES, case_workbook_path
from xlsx_support import add_table, load_case_excel, read_table, text_cell, workbook_bytes


CASE_HEADERS = ["用例编号", "模块路径", "用例名称", "优先级", "风险", "冒烟", "回归", "全量", "适用目标", "自动化评估", "前置条件", "测试数据", "操作步骤", "步骤预期", "测试点编号", "需求追溯", "运行条件编号", "影响", "需确认", "清理步骤", "标签", "备注", "变体来源"]
ASSERTION_HEADERS = ["用例编号", "断言编号", "步骤序号", "适用目标", "测试点编号", "检查内容", "观察位置", "预期结果", "所需证据", "比较方式", "比较值类型", "比较值", "检查时机", "判定标准"]
DERIVATION_HEADERS = ["规则编号", "来源片段或确认", "模型编号", "方法", "选择依据", "覆盖项编号", "维度", "条件", "操作", "预期", "准备方法", "方法数据", "处理", "排除理由", "测试点编号"]
REVIEW_HEADERS = ["阶段", "复查项目", "问题编号", "范围", "发现", "状态", "解决记录"]
AUTOMATION_HEADERS = ["用例编号", "目标编号", "可行性", "执行路线", "运行条件编号", "条件", "原因", "所需证据"]
RUNTIME_HEADERS = ["条件编号", "名称", "类型", "范围", "目标编号", "描述", "敏感", "提供方式", "备注"]
FEASIBILITY = {"automatable": "可自动化", "conditional": "有条件自动化", "manual": "人工执行"}
RISK = {"high": "高", "medium": "中", "low": "低"}
FORMAT_VERSION = "1.3"
COMPARISONS = {"equals":"相等", "contains":"包含文本", "number-equals":"数值相等", "number-delta":"数值变化", "unordered-equals":"列表相等(忽略顺序)", "unchanged":"前后不变", "evidence-review":"证据评判"}
VALUE_TYPES = ["文本", "数值", "布尔", "列表", "对象", "空值", "不适用"]


def comparison_cells(check):
    if "expected_value" not in check:
        return ["不适用", ""]
    value = check["expected_value"]
    if value is None:
        return ["空值", ""]
    if isinstance(value, bool):
        return ["布尔", "是" if value else "否"]
    if isinstance(value, str):
        return ["文本", value]
    kind = "数值" if isinstance(value, (int, float)) else "列表" if isinstance(value, list) else "对象"
    return [kind, json.dumps(value, ensure_ascii=False)]


def comparison_value(kind, text):
    if kind not in VALUE_TYPES:
        raise ValidationFailure(["unrecognized comparison value type"])
    if kind in {"不适用", "空值"}:
        if text:
            raise ValidationFailure([kind + ": comparison value must be empty"])
        return {} if kind == "不适用" else {"expected_value":None}
    if kind == "文本":
        return {"expected_value":text}
    if kind == "布尔":
        return {"expected_value":_yes(text, "comparison boolean")}
    value = json.loads(text)
    expected_type = {"数值":(int,float), "列表":list, "对象":dict}[kind]
    if isinstance(value, bool) or not isinstance(value, expected_type):
        raise ValidationFailure([kind + ": comparison value does not match its selected type"])
    return {"expected_value":value}


def _lines(values: list[Any]) -> str:
    return "\n".join(str(value) for value in values)


def _list(value: str) -> list[str]:
    return [line.strip() for line in value.replace("\r", "").split("\n") if line.strip()]


def _numbered(values: list[str]) -> str:
    return "\n".join(f"{i}. " + value.replace("\n", "\n   ") for i, value in enumerate(values, 1))


def _items(value: str, field: str) -> list[str]:
    return _parse_numbered(value, field) if value else []


def _parse_numbered(value: str, field: str) -> list[str]:
    matches = list(re.finditer(r"(?m)^(\d+)\.\s", value.replace("\r", "")))
    if not matches or value[:matches[0].start()].strip():
        raise ValidationFailure([f"{field}: use consecutive numbered lines such as 1. 操作"])
    if [int(m.group(1)) for m in matches] != list(range(1, len(matches) + 1)):
        raise ValidationFailure([f"{field}: numbering must start at 1 and be consecutive"])
    normalized = value.replace("\r", "")
    return [normalized[m.end():matches[i+1].start() if i+1 < len(matches) else len(normalized)].rstrip().replace("\n   ", "\n") for i, m in enumerate(matches)]


def _yes(value: str, field: str) -> bool:
    if value not in {"是", "否"}:
        raise ValidationFailure([f"{field}: expected 是 or 否"])
    return value == "是"


def _decode_label(value: str, labels: dict[str, str], field: str) -> str:
    reverse = {v: k for k, v in labels.items()}
    if value not in reverse:
        raise ValidationFailure([f"{field}: unrecognized value {value!r}"])
    return reverse[value]


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def module_paths(test_points: dict[str, Any] | None) -> dict[str, str]:
    result = {}
    def visit(node: dict[str, Any], path: list[str]) -> None:
        current = path + [node["title"]]
        if not node["children"]:
            result[node["id"]] = " / ".join(current[1:-1])
        for child in node["children"]:
            visit(child, current)
    if test_points:
        visit(test_points["tree"], [])
    return result


def render_case_workbook(manifest: dict[str, Any], manifest_sha256: str, paths: dict[str, str] | None, context: dict[str, Any], suite: str = "full") -> bytes:
    validate_manifest_data(manifest)
    from design_context import validate_design_context_data
    validate_design_context_data(context)
    paths = paths or {}
    workbook = Workbook()
    workbook.remove(workbook.active)
    case_rows, assertion_rows, automation_rows = [], [], []
    for case in manifest["cases"]:
        if suite not in case["suite_membership"]:
            continue
        summary = _lines([a["target_id"] + ": " + FEASIBILITY[a["feasibility"]] for a in case["automation"]])
        case_rows.append([
            case["id"], _lines(list(dict.fromkeys(paths.get(p, "") for p in case["test_point_refs"]))), case["title"], case["priority"], RISK[case["risk"]],
            *["是" if s in case["suite_membership"] else "否" for s in ("smoke", "regression", "full")],
            _lines(case["applicable_targets"]), summary, _numbered(case["preconditions"]), _numbered(case["test_data"]),
            _numbered([s["action"] for s in case["steps"]]), _numbered([s["expected"] for s in case["steps"]]),
            _lines(case["test_point_refs"]), _lines(case["source_refs"]), _lines(case["runtime_requirement_refs"]),
            case["safety"]["impact"], "是" if case["safety"]["requires_confirmation"] else "否", _numbered(case["safety"]["cleanup_steps"]),
            _numbered(case.get("tags", [])), case.get("notes", ""), case.get("variant_of", ""),
        ])
        for assertion in case["assertions"]:
            step_numbers = [i for i, s in enumerate(case["steps"], 1) if assertion["id"] in s["assertion_refs"]]
            assertion_rows.append([case["id"], assertion["id"], _lines(step_numbers), _lines(assertion["target_ids"]), _lines(assertion["test_point_refs"]), assertion["description"], assertion["observation"], assertion["expected"], _lines(assertion["required_evidence"])])
            check = assertion["check"]
            assertion_rows[-1].extend([COMPARISONS[check["kind"]], *comparison_cells(check), check["timing"], check.get("criteria", "")])
        for auto in case["automation"]:
            automation_rows.append([case["id"], auto["target_id"], FEASIBILITY[auto["feasibility"]], auto["candidate_route"], _lines(auto["runtime_requirement_refs"]), auto.get("condition", ""), auto.get("reason", ""), _lines(auto["required_evidence"])])
    add_table(workbook, "测试用例", CASE_HEADERS, case_rows,
              {"用例名称": 36, "模块路径": 32, "操作步骤": 55, "步骤预期": 55, "前置条件": 36, "测试数据": 36, "自动化评估": 27},
              {"冒烟": ["是", "否"], "回归": ["是", "否"], "全量": ["是", "否"], "需确认": ["是", "否"], "优先级": ["P0", "P1", "P2", "P3"], "风险": ["高", "中", "低"]})
    add_table(workbook, "必检断言", ASSERTION_HEADERS, assertion_rows, {"检查内容": 36, "观察位置": 40, "预期结果": 55, "检查时机": 40, "判定标准": 55, "比较值":40}, {"比较方式":list(COMPARISONS.values()), "比较值类型":VALUE_TYPES})
    add_table(workbook, "自动化评估", AUTOMATION_HEADERS, automation_rows, {"条件": 48, "原因": 48}, {"可行性": list(FEASIBILITY.values())})
    add_table(workbook, "运行条件", RUNTIME_HEADERS, [[r["id"], r["name"], r["kind"], r["scope"], _lines(r.get("target_ids", [])), r["description"], "是" if r["sensitive"] else "否", r["collection"], r.get("notes", "")] for r in manifest["runtime_requirements"]], {"描述": 60}, {"敏感": ["是", "否"], "范围": ["run", "target", "case"]})
    derivations, reviews = [], []
    for rule in context["rules"]:
        for model in rule["design_models"]:
            for item in model["items"]:
                method_data = {"model":{k:v for k,v in model.items() if k not in {"id","method","rationale","items"}}, "item":{k:v for k,v in item.items() if k not in {"id","dimensions","condition","action","expected","setup","disposition","reason","test_point_refs"}}}
                derivations.append([rule["id"], _lines(rule["segment_refs"] + rule["decision_refs"]), model["id"], model["method"], model["rationale"], item["id"], _lines(item["dimensions"]), item["condition"], item["action"], item["expected"], item["setup"], json.dumps(method_data, ensure_ascii=False), item["disposition"], item.get("reason",""), _lines(item["test_point_refs"])])
    for phase, review in [("需求与测试点", context["review"]), ("测试用例", manifest["review"])]:
        reviews += [[phase, check, "", "", "已复查", review["status"], review.get("notes", "")] for check in review["checks"]]
        reviews += [[phase, f["category"], f["id"], _lines(f["scope_refs"]), f["description"], f["status"], f.get("resolution", "")] for f in review["findings"]]
    add_table(workbook, "设计依据", DERIVATION_HEADERS, derivations, {"条件":40,"预期":50,"方法数据":65})
    add_table(workbook, "设计复查", REVIEW_HEADERS, reviews, {"发现":55,"解决记录":65})
    add_table(workbook, "使用说明", ["项目", "说明"], [
        ["项目", manifest["project"]["name"]],
        ["用例集", {"smoke":"冒烟", "regression":"回归", "full":"全量"}[suite] + "用例文件；共享同一基线编号。"],
        ["可编辑", "编辑用例业务字段、必检断言、各目标自动化评估和运行条件。模块路径和主表自动化摘要由基线生成，请在对应明细表修改。"],
        ["步骤", "操作步骤、步骤预期、前置条件、测试数据、清理步骤和标签采用 1.、2. 连续编号；条目内换行缩进三个空格。必检断言的步骤序号指向操作编号。多个编号或引用逐行填写。"],
        ["判定", "步骤预期供阅读，必检断言是执行判定依据。修改任一预期后必须重新审查二者及测试点的一致性。"],
        ["比较值", "文本和数值直接填写，布尔填是/否；仅列表和对象采用 JSON。空值及不适用留空。前后不变和证据评判的比较值类型选择不适用。检查时机必须明确，证据评判填写客观标准。"],
        ["设计依据与复查", "这两张表由已确认设计生成。发现遗漏时交给 Designer 更新规则、覆盖项和 XMind，再重新导出；直接修改会被拦截。"],
        ["修改交接", "修改 Excel 后使用 import-case-workbook 读回待审清单；完成语义审查后更新基线并重新导出。执行前 validate-case-workbook 必须通过。"],
        ["稳定编号", "保留已有用例、测试点和断言编号；新增或删除时同步各明细表。新业务规则需要更新并确认 XMind。"],
        ["运行条件", "此表仅定义需要什么；实际环境和安全会话引用在运行时收集。请勿填密码、令牌或登录 cookie。"],
        ["单元格", "长内容完整保存在单元格中，可在编辑栏查看；禁止在业务单元格使用公式。"],
    ], {"项目": 24, "说明": 110})
    for sheet in ("设计依据", "设计复查", "运行条件"):
        workbook[sheet].sheet_state = "hidden"
    metadata = workbook.create_sheet("_baseline")
    metadata.sheet_state = "veryHidden"
    payload = _canonical(manifest)
    path_payload = _canonical(paths)
    design_payload = _canonical({"设计依据":derivations, "设计复查":reviews})
    metadata_rows = [["format", FORMAT_VERSION], ["suite", suite], ["manifest_sha256", manifest_sha256], ["payload_sha256", hashlib.sha256(payload.encode()).hexdigest()], ["paths_sha256", hashlib.sha256(path_payload.encode()).hexdigest()], ["design_sha256", hashlib.sha256(design_payload.encode()).hexdigest()]]
    # JSON chunks retain non-editable source and project metadata losslessly.
    for i in range(0, len(payload), 30000):
        metadata_rows.append(["payload", payload[i:i+30000]])
    for i in range(0, len(path_payload), 30000):
        metadata_rows.append(["paths", path_payload[i:i+30000]])
    for i in range(0, len(design_payload), 30000):
        metadata_rows.append(["design", design_payload[i:i+30000]])
    for r, values in enumerate(metadata_rows, 1):
        for c, value in enumerate(values, 1):
            text_cell(metadata.cell(r, c), value)
    return workbook_bytes(workbook)


def read_case_workbook(path: Path) -> tuple[dict[str, Any], dict[str, Any], str]:
    workbook = load_case_excel(path)
    try:
        expected_sheets = {"测试用例", "必检断言", "自动化评估", "运行条件", "设计依据", "设计复查", "使用说明", "_baseline"}
        if set(workbook.sheetnames) != expected_sheets:
            raise ValidationFailure(["Excel worksheets changed; preserve the template and its baseline metadata"])
        metadata_sheet = workbook["_baseline"]
        metadata_sheet.reset_dimensions()
        metadata = list(metadata_sheet.iter_rows(values_only=True))
        values = {str(r[0]): str(r[1]) for r in metadata if r[0] not in {"payload", "paths", "design"}}
        if values.get("format") != FORMAT_VERSION:
            raise ValidationFailure(["unsupported case workbook format; regenerate with the current runtime"])
        payload = "".join(str(r[1]) for r in metadata if r[0] == "payload")
        if hashlib.sha256(payload.encode()).hexdigest() != values.get("payload_sha256"):
            raise ValidationFailure(["Excel baseline metadata was modified"])
        original = json.loads(payload)
        suite = values.get("suite")
        if suite not in SUITES:
            raise ValidationFailure(["invalid workbook suite"])
        visible_ids = {c["id"] for c in original["cases"] if suite in c["suite_membership"]}
        validate_manifest_data(original)
        path_payload = "".join(str(r[1]) for r in metadata if r[0] == "paths")
        if hashlib.sha256(path_payload.encode()).hexdigest() != values.get("paths_sha256"):
            raise ValidationFailure(["Excel module path metadata was modified"])
        paths = json.loads(path_payload)
        design_payload = "".join(str(r[1]) for r in metadata if r[0] == "design")
        if hashlib.sha256(design_payload.encode()).hexdigest() != values.get("design_sha256"):
            raise ValidationFailure(["Excel design metadata was modified"])
        generated = json.loads(design_payload)
        for sheet, headers in [("设计依据", DERIVATION_HEADERS), ("设计复查", REVIEW_HEADERS)]:
            actual_rows = [[row[h] for h in headers] for row in read_table(workbook, sheet, headers)]
            if sorted(map(_canonical, actual_rows)) != sorted(map(_canonical, generated[sheet])):
                raise ValidationFailure([sheet + ": generated design rows changed; update Designer baselines and rerender"])
        candidate = copy.deepcopy(original)
        candidate["cases"] = []
        case_by_id = {}
        for row in read_table(workbook, "测试用例", CASE_HEADERS):
            cid = row["用例编号"]
            if any(c["id"] == cid for c in original["cases"]) and cid not in visible_ids:
                raise ValidationFailure([f"{cid}: belongs to another suite view"])
            if cid in case_by_id:
                raise ValidationFailure([f"duplicate Excel case id: {cid}"])
            actions = _parse_numbered(row["操作步骤"], cid + "/操作步骤")
            expected = _parse_numbered(row["步骤预期"], cid + "/步骤预期")
            if len(actions) != len(expected):
                raise ValidationFailure([f"{cid}: action and expected step counts differ"])
            refs = _list(row["测试点编号"])
            derived_path = _lines(list(dict.fromkeys(paths.get(p, "") for p in refs)))
            if row["模块路径"] != derived_path:
                raise ValidationFailure([f"{cid}: module path is generated; update test-point references and rerender"])
            case = {
                "id": cid, "title": row["用例名称"], "source_refs": _list(row["需求追溯"]), "test_point_refs": refs,
                "priority": row["优先级"], "risk": _decode_label(row["风险"], RISK, cid),
                "suite_membership": [s for s, label in [("smoke", "冒烟"), ("regression", "回归"), ("full", "全量")] if _yes(row[label], cid + "/" + label)],
                "applicable_targets": _list(row["适用目标"]), "runtime_requirement_refs": _list(row["运行条件编号"]),
                "preconditions": _items(row["前置条件"], cid + "/前置条件"), "test_data": _items(row["测试数据"], cid + "/测试数据"),
                "steps": [{"action": a, "expected": e, "assertion_refs": []} for a, e in zip(actions, expected)],
                "assertions": [], "automation": [], "safety": {"impact": row["影响"], "requires_confirmation": _yes(row["需确认"], cid), "cleanup_steps": _items(row["清理步骤"], cid + "/清理步骤")},
            }
            for field, label in [("tags", "标签"), ("notes", "备注"), ("variant_of", "变体来源")]:
                if row[label]:
                    case[field] = _items(row[label], cid + "/标签") if field == "tags" else row[label]
            old_case = next((c for c in original["cases"] if c["id"] == cid), None)
            if old_case:
                for field, label in (("preconditions", "前置条件"), ("test_data", "测试数据")):
                    if row[label] == _numbered(old_case[field]):
                        case[field] = copy.deepcopy(old_case[field])
                if row["清理步骤"] == _numbered(old_case["safety"]["cleanup_steps"]):
                    case["safety"]["cleanup_steps"] = copy.deepcopy(old_case["safety"]["cleanup_steps"])
                for field, label in (("action", "操作步骤"), ("expected", "步骤预期")):
                    if row[label] == _numbered([s[field] for s in old_case["steps"]]):
                        for step, old_step in zip(case["steps"], old_case["steps"]):
                            step[field] = old_step[field]
                if row["标签"] == _numbered(old_case.get("tags", [])) and "tags" in old_case:
                    case["tags"] = copy.deepcopy(old_case["tags"])
            if old_case and old_case["safety"].get("notes"):
                case["safety"]["notes"] = old_case["safety"]["notes"]
            # Retain intentionally empty optional fields for a lossless unchanged round trip.
            if old_case:
                for field in ("tags", "notes", "variant_of"):
                    if field in old_case and field not in case:
                        case[field] = [] if field == "tags" else ""
            candidate["cases"].append(case)
            case_by_id[cid] = (case, row)
        for row in read_table(workbook, "必检断言", ASSERTION_HEADERS):
            cid = row["用例编号"]
            if cid not in case_by_id:
                raise ValidationFailure([f"assertion references unknown case {cid}"])
            case = case_by_id[cid][0]
            assertion = {"id": row["断言编号"], "description": row["检查内容"], "observation": row["观察位置"], "expected": row["预期结果"], "target_ids": _list(row["适用目标"]), "test_point_refs": _list(row["测试点编号"]), "required_evidence": _list(row["所需证据"])}
            check = {"kind":_decode_label(row["比较方式"], COMPARISONS, assertion["id"]), "timing":row["检查时机"]}
            check.update(comparison_value(row["比较值类型"], row["比较值"]))
            if row["判定标准"]:
                check["criteria"] = row["判定标准"]
            old_assertion = next((a for c in original["cases"] if c["id"] == cid for a in c["assertions"] if a["id"] == assertion["id"]), {})
            if "criteria" in old_assertion.get("check", {}) and "criteria" not in check:
                check["criteria"] = ""
            assertion["check"] = check
            case["assertions"].append(assertion)
            for text in _list(row["步骤序号"]):
                try:
                    number = int(text)
                except ValueError as exc:
                    raise ValidationFailure([f"{cid}/{assertion['id']}: invalid step number {text!r}"]) from exc
                if number < 1 or number > len(case["steps"]):
                    raise ValidationFailure([f"{cid}/{assertion['id']}: step number is out of range"])
                case["steps"][number - 1]["assertion_refs"].append(assertion["id"])
        for row in read_table(workbook, "自动化评估", AUTOMATION_HEADERS):
            cid = row["用例编号"]
            if cid not in case_by_id:
                raise ValidationFailure([f"automation references unknown case {cid}"])
            auto = {"target_id": row["目标编号"], "feasibility": _decode_label(row["可行性"], FEASIBILITY, cid), "candidate_route": row["执行路线"], "runtime_requirement_refs": _list(row["运行条件编号"]), "required_evidence": _list(row["所需证据"])}
            for field, label in [("condition", "条件"), ("reason", "原因")]:
                if row[label]:
                    auto[field] = row[label]
                else:
                    old = next((a for c in original["cases"] if c["id"] == cid for a in c["automation"] if a["target_id"] == auto["target_id"]), {})
                    if field in old:
                        auto[field] = ""
            case_by_id[cid][0]["automation"].append(auto)
        for cid, (case, row) in case_by_id.items():
            original_case = next((c for c in original["cases"] if c["id"] == cid), None)
            original_summary = _lines([a["target_id"] + ": " + FEASIBILITY[a["feasibility"]] for a in original_case["automation"]]) if original_case else ""
            if row["自动化评估"] != original_summary:
                raise ValidationFailure([f"{cid}: automation summary is generated; edit 自动化评估 worksheet"])
        candidate["runtime_requirements"] = []
        for row in read_table(workbook, "运行条件", RUNTIME_HEADERS):
            requirement = {"id": row["条件编号"], "name": row["名称"], "kind": row["类型"], "scope": row["范围"], "description": row["描述"], "sensitive": _yes(row["敏感"], row["条件编号"]), "collection": row["提供方式"]}
            if row["目标编号"]:
                requirement["target_ids"] = _list(row["目标编号"])
            if row["备注"]:
                requirement["notes"] = row["备注"]
            old = next((r for r in original["runtime_requirements"] if r["id"] == requirement["id"]), {})
            for field in ("target_ids", "notes"):
                if field in old and field not in requirement:
                    requirement[field] = [] if field == "target_ids" else ""
            candidate["runtime_requirements"].append(requirement)
        candidate["cases"].extend(copy.deepcopy(c) for c in original["cases"] if c["id"] not in visible_ids)
        # Sorting rows in Excel is a view change, not a requirement change.
        def baseline_order(items, baseline, key):
            order = {item[key]: i for i, item in enumerate(baseline)}
            items.sort(key=lambda item: (order.get(item[key], len(order)), item[key]))
        baseline_order(candidate["cases"], original["cases"], "id")
        baseline_order(candidate["runtime_requirements"], original["runtime_requirements"], "id")
        for case in candidate["cases"]:
            old = next((c for c in original["cases"] if c["id"] == case["id"]), None)
            if old:
                baseline_order(case["automation"], old["automation"], "target_id")
                baseline_order(case["assertions"], old["assertions"], "id")
                for i, step in enumerate(case["steps"]):
                    if i < len(old["steps"]) and set(step["assertion_refs"]) == set(old["steps"][i]["assertion_refs"]):
                        step["assertion_refs"] = old["steps"][i]["assertion_refs"][:]
        changed = _canonical(candidate) != _canonical(original)
        if changed:
            candidate["review"]["status"] = "pending"
            candidate["review"].pop("reviewed_at", None)
            candidate["review"]["checks"] = []
            candidate["review"]["notes"] = "Imported Excel edits require source, test-point, step, and assertion consistency review. Previous findings are retained; repeat checks for the new baseline."
        validate_manifest_data(candidate, final=not changed)
        return candidate, original, values["manifest_sha256"]
    except ValidationFailure:
        raise
    except (ValueError, KeyError, TypeError, IndexError) as exc:
        raise ValidationFailure([f"invalid case workbook: {type(exc).__name__}: {exc}"]) from exc
    finally:
        workbook.close()


def validate_case_workbook_file(workbook_path: Path, manifest_path: Path) -> None:
    manifest = validate_manifest_file(manifest_path)
    candidate, original, digest = read_case_workbook(workbook_path)
    if digest != sha256_file(manifest_path) or _canonical(original) != _canonical(manifest):
        raise ValidationFailure(["Excel workbook belongs to a stale or different manifest; reconcile and rerender"])
    if _canonical(candidate) != _canonical(manifest):
        raise ValidationFailure(["Excel has unreviewed edits; import-case-workbook, review, and rerender before execution"])


def validate_case_workbooks(manifest_path):
    for suite in SUITES:
        validate_case_workbook_file(case_workbook_path(manifest_path, suite), manifest_path)
        book = load_case_excel(case_workbook_path(manifest_path, suite))
        try:
            rows = list(book["_baseline"].iter_rows(values_only=True))
            if dict((r[0], r[1]) for r in rows if r[0] == "suite").get("suite") != suite:
                raise ValidationFailure(["workbook filename and suite metadata disagree"])
        finally:
            book.close()


def render_case_workbooks(manifest_path):
    from artifact_tools import validate_test_points_file
    from design_context import validate_design_context_file
    manifest = validate_manifest_file(manifest_path)
    paths = module_paths(validate_test_points_file(manifest_path.with_name("test-points.json")))
    context = validate_design_context_file(manifest_path.with_name("design-context.json"))
    outputs = []
    for suite in SUITES:
        output = case_workbook_path(manifest_path, suite)
        # Never discard edits to a workbook from this or an earlier baseline.
        if output.exists():
            candidate, original, _ = read_case_workbook(output)
            if _canonical(candidate) != _canonical(original):
                for field in ("cases", "runtime_requirements"):
                    before = {x["id"]:x for x in original[field]}
                    after = {x["id"]:x for x in candidate[field]}
                    reviewed = {x["id"]:x for x in manifest[field]}
                    if any(before.get(k) != after.get(k) and reviewed.get(k) != after.get(k) for k in set(before) | set(after)):
                        raise ValidationFailure([str(output) + ": unreviewed edits; import and reconcile before exporting"])
        outputs.append((output, render_case_workbook(manifest, sha256_file(manifest_path), paths, context, suite)))
    for output, data in outputs:
        write_bytes_atomic(output, data)
    validate_case_workbooks(manifest_path)
    return [p for p, _ in outputs]


def import_case_workbooks(manifest_path):
    manifest = validate_manifest_file(manifest_path)
    merged = copy.deepcopy(manifest)
    changes = {}
    for suite in SUITES:
        candidate, original, digest = read_case_workbook(case_workbook_path(manifest_path, suite))
        if digest != sha256_file(manifest_path) or _canonical(original) != _canonical(manifest):
            raise ValidationFailure(["stale suite workbook; reconcile its baseline before import"])
        for field, key in [("cases", "id"), ("runtime_requirements", "id")]:
            before = {x[key]:x for x in original[field]}
            after = {x[key]:x for x in candidate[field]}
            for ident in set(before) | set(after):
                if before.get(ident) == after.get(ident):
                    continue
                marker = (field, ident)
                value = after.get(ident)
                if marker in changes and changes[marker] != value:
                    raise ValidationFailure([f"conflicting edits across suite workbooks: {ident}"])
                changes[marker] = value
    for (field, ident), value in changes.items():
        merged[field] = [x for x in merged[field] if x["id"] != ident]
        if value is not None:
            merged[field].append(value)
    if changes:
        merged["review"].update(status="pending", checks=[], notes="Suite workbook edits merged; semantic review is required.")
        merged["review"].pop("reviewed_at", None)
    validate_manifest_data(merged, final=not changes)
    return merged


def main() -> int:
    parser = argparse.ArgumentParser(description="Export three suite files, validate all views, or reconcile Excel edits.")
    parser.add_argument("operation", choices=("render", "import", "validate"))
    parser.add_argument("manifest", type=Path)
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    manifest_path = args.manifest.resolve()
    try:
        if args.operation == "render":
            if args.output or args.workbook:
                raise ValidationFailure(["render exports the three canonical suite files; no single-workbook output"])
            for output in render_case_workbooks(manifest_path):
                print(f"WROTE: {output}")
        elif args.operation == "validate":
            validate_case_workbooks(manifest_path)
            print("OK: all three Excel suite views match the reviewed manifest")
        else:
            manifest = validate_manifest_file(manifest_path)
            if args.workbook:
                candidate, original, digest = read_case_workbook(args.workbook.resolve())
                if digest != sha256_file(manifest_path) or _canonical(original) != _canonical(manifest):
                    raise ValidationFailure(["stale workbook baseline"])
            else:
                candidate = import_case_workbooks(manifest_path)
            output = (args.output or manifest_path.with_name("test-manifest.pending.json")).resolve()
            if output in {manifest_path, manifest_path.with_name("test-points.json"), manifest_path.with_name("design-context.json")} or output.suffix != ".json":
                raise ValidationFailure(["import must write separate pending JSON state"])
            write_text_atomic(output, json.dumps(candidate, ensure_ascii=False, indent=2) + "\n")
            print(f"WROTE: {output}; review={candidate['review']['status']}")
    except (ValidationFailure, OSError) as exc:
        for error in exc.errors if isinstance(exc, ValidationFailure) else [str(exc)]:
            print(f"ERROR: {error}")
        return 1
    return 0


def render_main() -> int:
    import sys
    sys.argv.insert(1, "render")
    return main()


def import_main() -> int:
    import sys
    sys.argv.insert(1, "import")
    return main()


def validate_main() -> int:
    import sys
    sys.argv.insert(1, "validate")
    return main()

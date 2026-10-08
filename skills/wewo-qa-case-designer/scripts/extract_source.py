from __future__ import annotations

"""Inventory local source content; visual interpretation remains a host capability."""

import argparse
import hashlib
import json
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree as ET

from artifact_tools import ValidationFailure, sha256_file, write_bytes_atomic, write_text_atomic


def xml(data: bytes):
    if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
        raise ValidationFailure(["XML entity declarations are not supported"])
    return ET.fromstring(data)


def local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


class HTMLContent(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts, self.links, self.images = [], [], []
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        values = dict(attrs)
        if tag in {"script", "style"}:
            self.hidden += 1
        if tag in {"p", "br", "div", "tr", "li", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")
        if tag == "a" and values.get("href"):
            self.links.append(values["href"])
        if tag == "img":
            self.images.append(values)

    def handle_endtag(self, tag):
        if tag in {"script", "style"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


class Extractor:
    def __init__(self, source: Path, source_id: str, output: Path):
        self.source, self.source_id, self.output = source, source_id, output
        self.segments, self.warnings = [], []

    def add(self, anchor: str, kind: str, content: str, *, asset=None, review=False):
        if not content.strip() and not asset:
            return
        segment = {
            "id": f"SEG-{self.source_id.removeprefix('SRC-')}-{len(self.segments)+1:05d}",
            "source_id": self.source_id, "anchor": anchor, "kind": kind,
            "content": content, "read_status": "needs-review" if review else "read",
        }
        if asset:
            segment["asset_paths"] = [asset]
        self.segments.append(segment)

    def asset(self, data: bytes, name: str) -> str:
        suffix = Path(name).suffix.lower()
        # Content addressing avoids trusting archive filenames or overwriting source files.
        filename = hashlib.sha256(data).hexdigest() + suffix
        path = self.output.parent / (self.output.stem + ".assets") / filename
        write_bytes_atomic(path, data)
        return path.relative_to(self.output.parent).as_posix()

    def archive_assets(self, archive):
        for name in archive.namelist():
            if Path(name).suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".svg", ".emf", ".wmf", ".tif", ".tiff"}:
                self.add("asset/" + name, "image", "Embedded image; inspect in its original context", asset=self.asset(archive.read(name), name), review=True)

    def xmind(self, archive):
        if "content.json" in archive.namelist():
            sheets = json.loads(archive.read("content.json"))
            def topic(node, prefix):
                anchor = prefix + "/topic:" + str(node.get("id", len(self.segments)))
                title = str(node.get("title", ""))
                self.add(anchor, "topic", title or "Untitled topic; inspect its context", review=not bool(title.strip()))
                for key in ("notes", "labels", "markers", "href", "image", "boundaries", "summaries", "extensions"):
                    if node.get(key):
                        self.add(anchor + "/" + key, "link" if key == "href" else "note", json.dumps(node[key], ensure_ascii=False))
                children = node.get("children", {})
                for group, nodes in children.items():
                    for child in nodes:
                        topic(child, anchor + "/" + group)
            for i, sheet in enumerate(sheets, 1):
                prefix = f"sheet:{i}"
                self.add(prefix, "sheet", sheet.get("title", ""))
                topic(sheet["rootTopic"], prefix)
                for key in ("relationships", "notes", "extensions"):
                    if sheet.get(key):
                        self.add(prefix + "/" + key, "diagram", json.dumps(sheet[key], ensure_ascii=False))
        elif "content.xml" in archive.namelist():
            root = xml(archive.read("content.xml"))
            def visit(element, prefix):
                tag = local(element.tag)
                anchor = prefix + "/" + tag + ":" + str(element.get("id", len(self.segments)))
                if tag == "topic":
                    title = next(("".join(c.itertext()) for c in element if local(c.tag) == "title"), "")
                    self.add(anchor, "topic", title or "Untitled topic; inspect its context", review=not bool(title.strip()))
                if tag in {"notes", "labels", "marker-refs", "relationship", "boundary", "summary", "image", "extensions"}:
                    self.add(anchor, "diagram" if tag in {"relationship", "boundary", "summary"} else "note", ET.tostring(element, encoding="unicode"))
                    return
                href = next((v for k, v in element.attrib.items() if local(k) == "href"), None)
                if href:
                    self.add(anchor + "/href", "link", href)
                for child in element:
                    visit(child, anchor)
            visit(root, "xmind")
        else:
            raise ValidationFailure(["XMind has no supported content.xml/content.json"])
        self.archive_assets(archive)

    def office(self, archive, kind):
        if kind == "docx":
            from docx import Document
            document = Document(self.source)
            for i, table in enumerate(document.tables, 1):
                for j, row in enumerate(table.rows, 1):
                    self.add(f"word/table:{i}/row:{j}", "table", "\t".join(f"col:{k}={cell.text}" for k, cell in enumerate(row.cells, 1)))
        else:
            from pptx import Presentation
            presentation = Presentation(self.source)
            for i, slide in enumerate(presentation.slides, 1):
                self.add(f"slide:{i}/layout", "slide", f"Slide {i}: {len(slide.shapes)} shapes; inspect their spatial relationships", review=True)
                for shape in slide.shapes:
                    if shape.has_table:
                        for j, row in enumerate(shape.table.rows, 1):
                            self.add(f"slide:{i}/shape:{shape.shape_id}/row:{j}", "table", "\t".join(f"col:{k}={cell.text}" for k, cell in enumerate(row.cells, 1)))
        # OOXML parts are parsed to include text boxes, notes, comments and headers,
        # which high-level paragraph-only readers frequently omit.
        prefixes = {"docx": ("word/",), "pptx": ("ppt/",)}[kind]
        for name in archive.namelist():
            if not name.startswith(prefixes) or not name.endswith(".xml"):
                continue
            root = xml(archive.read(name))
            blocks = [e for e in root.iter() if local(e.tag) in {"p", "comment"}]
            for i, block in enumerate(blocks, 1):
                content = "".join((e.text or "") if local(e.tag) in {"t", "instrText"} else "\t" if local(e.tag) == "tab" else "\n" if local(e.tag) in {"br", "cr"} else "" for e in block.iter())
                self.add(f"{name}/block:{i}", "note" if "notes" in name or "comments" in name else "text", content)
            visuals = [e for e in root.iter() if local(e.tag) in {"drawing", "pict", "graphicFrame", "cxnSp"}]
            for i, visual in enumerate(visuals, 1):
                self.add(f"{name}/visual:{i}", "diagram", ET.tostring(visual, encoding="unicode"), review=True)
        for name in archive.namelist():
            if name.endswith(".rels"):
                for i, relation in enumerate(xml(archive.read(name)), 1):
                    if relation.get("TargetMode") == "External":
                        self.add(f"{name}/external:{i}", "link", relation.get("Target", ""))
        self.archive_assets(archive)
        self.warnings.append("OOXML text preserves part/block anchors; inspect the original layout for table relationships, diagrams and floating objects.")
        self.add("document-layout", "other", "Review original Office layout, including tables, merged cells and shape relationships", review=True)

    def excel(self):
        from openpyxl import load_workbook
        # Non-streaming reads keep comments, merge ranges and drawings available.
        workbook = load_workbook(self.source, data_only=False, keep_links=False)
        try:
            for sheet in workbook:
                for row in sheet:
                    cells = [f"{cell.coordinate}={cell.value}" for cell in row if cell.value is not None]
                    if cells:
                        self.add(f"sheet:{sheet.title}/row:{row[0].row}", "table", "\t".join(cells), review=any(cell.data_type == "f" for cell in row))
                    for cell in row:
                        if cell.comment:
                            self.add(f"sheet:{sheet.title}/comment:{cell.coordinate}", "note", cell.comment.text)
                if sheet.merged_cells.ranges:
                    self.add(f"sheet:{sheet.title}/merges", "note", ", ".join(str(r) for r in sheet.merged_cells.ranges))
                if sheet.sheet_state != "visible":
                    self.add(f"sheet:{sheet.title}/visibility", "note", sheet.sheet_state)
            with zipfile.ZipFile(self.source) as archive:
                self.archive_assets(archive)
                if any(name.startswith("xl/drawings/") for name in archive.namelist()):
                    self.add("workbook-drawings", "diagram", "Inspect workbook drawings, charts and their relationships", review=True)
            self.warnings.append("Formula cells are preserved as formulas; inspect original calculated values if they express a requirement.")
        finally:
            workbook.close()

    def pdf(self):
        from pypdf import PdfReader
        reader = PdfReader(self.source)
        if reader.is_encrypted:
            raise ValidationFailure(["encrypted PDF requires an authorized decrypted export"])
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            self.add(f"page:{i}/text", "text", text)
            self.add(f"page:{i}/visual", "page", "Render and inspect original page for scanned text, tables, layout and diagrams", review=True)
            for j, picture in enumerate(page.images, 1):
                self.add(f"page:{i}/image:{j}", "image", "Embedded PDF image", asset=self.asset(picture.data, picture.name), review=True)
        self.warnings.append("Text extraction cannot certify visual completeness; every PDF page requires visual review.")

    def run(self):
        suffix = self.source.suffix.lower()
        with self.source.open("rb") as stream:
            signature = stream.read(5)
        if zipfile.is_zipfile(self.source):
            with zipfile.ZipFile(self.source) as archive:
                if sum(info.file_size for info in archive.infolist()) > 256 * 1024 * 1024 or len(archive.infolist()) > 20000:
                    raise ValidationFailure(["archive exceeds extraction limits; split the authorized source"])
                names = archive.namelist()
                if "content.json" in names or "content.xml" in names:
                    self.xmind(archive)
                elif "xl/workbook.xml" in names:
                    self.excel()
                elif "word/document.xml" in names:
                    self.office(archive, "docx")
                elif "ppt/presentation.xml" in names:
                    self.office(archive, "pptx")
                else:
                    raise ValidationFailure(["unsupported archive; provide a readable document export"])
        elif signature == b"%PDF-":
            self.pdf()
        elif suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}:
            from PIL import Image
            with Image.open(self.source) as picture:
                self.add("image:1", "image", f"{picture.format} {picture.width}x{picture.height}; visually transcribe text and relationships", asset=self.asset(self.source.read_bytes(), self.source.name), review=True)
        elif suffix in {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm"}:
            data = self.source.read_bytes()
            try:
                content = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                try:
                    content = data.decode("gb18030")
                    self.warnings.append("Decoded using GB18030; confirm original character encoding.")
                except UnicodeDecodeError as exc:
                    raise ValidationFailure(["cannot decode text; provide its encoding or a UTF-8 export"]) from exc
            if suffix in {".html", ".htm"}:
                parser = HTMLContent()
                parser.feed(content)
                self.add("html/text", "text", "".join(parser.parts))
                for i, link in enumerate(parser.links, 1):
                    self.add(f"html/link:{i}", "link", link)
                for i, picture in enumerate(parser.images, 1):
                    self.add(f"html/image:{i}", "image", json.dumps(picture, ensure_ascii=False), review=True)
                self.add("html/layout", "page", "Inspect rendered HTML and linked requirement pages", review=True)
            else:
                # Line anchors remain stable without discarding blank-line structure.
                for i, line in enumerate(content.splitlines(), 1):
                    self.add(f"line:{i}", "text", line)
        else:
            raise ValidationFailure([f"unsupported source type {suffix}; use a host reader or request a readable export"])
        if not self.segments:
            raise ValidationFailure(["source contains no extractable content; inspect the original before proceeding"])
        return {
            "format_version": "1.0", "source_id": self.source_id,
            "path": str(self.source), "sha256": sha256_file(self.source),
            "inventory_status": "partial" if any(s["read_status"] != "read" for s in self.segments) else "complete",
            "method": "bundled structural extraction", "segments": self.segments, "warnings": self.warnings,
        }


def main():
    parser = argparse.ArgumentParser(description="Extract a local requirement source without claiming visual interpretation.")
    parser.add_argument("source", type=Path)
    parser.add_argument("--source-id", required=True)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    if not re.fullmatch(r"SRC-[0-9]{3,}", args.source_id):
        parser.error("source-id must use SRC- followed by at least three digits")
    try:
        source, output = args.source.resolve(), args.output.resolve()
        if source == output:
            raise ValidationFailure(["output must not overwrite the source"])
        result = Extractor(source, args.source_id, output).run()
        write_text_atomic(output, json.dumps(result, ensure_ascii=False, indent=2) + "\n")
    except (ValidationFailure, OSError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        print(f"ERROR: {exc}")
        return 1
    print(f"WROTE: {output}; inventory={result['inventory_status']}; segments={len(result['segments'])}")
    return 0

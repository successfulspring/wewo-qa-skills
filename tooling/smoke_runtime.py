from __future__ import annotations

"""Maintenance check: use the bundled executable for complete artifact operations."""

import argparse
import json
import copy
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(ROOT / "tooling/runtime"), str(ROOT / "skills/wewo-qa-case-designer/scripts"), str(ROOT / "skills/wewo-qa-case-executor/scripts"), str(ROOT / "tests")]
from artifact_factory import artifact_set, execution_set, write_json, mutate_observation, seal_synthetic_points, seal_synthetic_readiness
from artifact_tools import sha256_file
from openpyxl import load_workbook
from runtime_version import VERSION


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--runtime", required=True, type=Path)
    args = parser.parse_args()
    runtime = args.runtime.resolve()
    def run(*arguments, fail=False):
        result = subprocess.run([str(runtime), *map(str, arguments)], text=True, encoding="utf-8", errors="replace", capture_output=True, timeout=90)
        if (result.returncode != 0) != fail:
            raise RuntimeError(f"command {' '.join(map(str, arguments))}: {result.returncode}\n{result.stdout}\n{result.stderr}")
        return result.stdout
    if run("--version").strip() != f"wewo-qa {VERSION}":
        raise RuntimeError("runtime version does not match source")
    # Reuse the maintenance protocol fixture to exercise real process launch,
    # collection and complete-run rendering through the packaged executable.
    # Its observations are synthetic; this is not a product acceptance test.
    from test_workflow import NativeRunnerTests
    fixture = NativeRunnerTests()
    fixture.setUp()
    try:
        plan_path = fixture.profile_path.with_name("native-plan.json")
        write_json(plan_path, fixture.plan)
        write_json(fixture.repo / "fixture-events.json", {
            "events": fixture.events,
            "report": '<testsuites><testsuite><testcase classname="Synthetic" name="acceptance"/></testsuite></testsuites>',
            "exit_code": 0,
        })
        run("run-native-tests", fixture.path, fixture.profile_path, plan_path)
        result_path = next(fixture.profile_path.parent.glob("execution-results.*.json"))
        run("validate-execution-results", result_path, fixture.path)
        run("judge-execution-results", result_path, fixture.path)
        judged_path = result_path.with_name(result_path.stem + ".judged.json")
        run("validate-execution-results", judged_path, fixture.path)
        run("render-execution-report", result_path, fixture.path, fail=True)
        run("summarize-execution", fixture.path, fixture.profile_path)
        summary_path = fixture.profile_path.with_name("execution-summary.json")
        run("render-execution-report", summary_path, fixture.path)
        book = load_workbook(fixture.profile_path.with_name("test-execution-report.xlsx"))
        if book["执行汇总"]["B3"].value != "PASSED" or book["执行历史"].max_row != 2:
            raise RuntimeError("native invocation was not preserved in the complete run report")
        book.close()
    finally:
        fixture.doCleanups()
    with tempfile.TemporaryDirectory(prefix="wewo-qa-smoke-") as folder:
        directory = Path(folder)
        manifest, path, profile, profile_path, results, results_path = execution_set(directory)
        # Exercise the native plan command with a genuine pre-leaf synthetic plan.
        original_context = json.loads((directory / "design-context.json").read_text(encoding="utf-8"))
        planned = copy.deepcopy(original_context)
        planned.update(coverage_plan={"status":"draft"}, review={"status":"draft"})
        for rule in planned["rules"]:
            for dimension in rule["coverage"]:
                dimension["test_point_refs"] = []
            for model in rule["design_models"]:
                for item in model["items"]:
                    item["test_point_refs"] = []
        planned_path = directory / "planned-context.json"
        write_json(planned_path, planned)
        run("validate-requirement-readiness", planned_path)
        run("validate-requirement-readiness", planned_path, "--confirmed")
        # Same decision, a new question occurrence and a free-form reply. Old
        # answers must not close a follow-up, even with refreshed fixture hashes.
        dialogue = copy.deepcopy(planned)
        dialogue["dialogue_events"] = dialogue["dialogue_events"][:2]
        dialogue["requirement_confirmation"] = {"status":"draft"}
        dialogue["dialogue_events"][0]["question_ids"] = ["Q1.1"]
        dialogue["dialogue_events"][0]["text"] = "[SYNTHETIC] Q1.1 What are the failure threshold and lock duration?"
        dialogue["decisions"][0]["question_id"] = "Q1.1"
        dialogue["dialogue_events"].append({"id":"EVT-QUESTION-002", "kind":"question", "message_ref":"synthetic://q2",
            "text":"[SYNTHETIC] Q1.1 Please restate that rule without choosing an option.", "question_ids":["Q1.1"],
            "revisits":{"Q1.1":"EVT-QUESTION-001"}})
        dialogue_path = directory / "followup-context.json"
        seal_synthetic_readiness(dialogue)
        write_json(dialogue_path, dialogue)
        run("validate-requirement-readiness", dialogue_path, fail=True)
        dialogue["dialogue_events"].append({"id":"EVT-ANSWER-002", "kind":"answer", "message_ref":"synthetic://a2",
            "text":"[SYNTHETIC] Five failures lock the account for fifteen minutes.", "question_ids":["Q1.1"], "in_reply_to":"EVT-QUESTION-002"})
        dialogue["branch_reviews"].append({"answer_event_ref":"EVT-ANSWER-002", "unit_refs":["RU-002"],
            "note":"Synthetic restatement preserves the same threshold and duration; no new branch."})
        old = dialogue["decisions"][0]
        replacement = {**old, "id":"DEC-002", "response_event_ref":"EVT-ANSWER-002"}
        replacement.pop("selected_options", None)
        old.update(status="superseded", superseded_by="DEC-002")
        dialogue["decisions"].append(replacement)
        for rule in dialogue["rules"]:
            rule["decision_refs"] = ["DEC-002" if ref == "DEC-001" else ref for ref in rule["decision_refs"]]
        for item in dialogue["requirement_items"]:
            if item["origin_ref"] == "DEC-001":
                item["origin_ref"] = "DEC-002"
        seal_synthetic_readiness(dialogue)
        write_json(dialogue_path, dialogue)
        run("validate-requirement-readiness", dialogue_path)
        run("validate-requirement-readiness", dialogue_path, "--confirmed", fail=True)
        # The extracted runtime must contain the language catalog and retain
        # typed Chinese literals while generating and importing English views.
        english_dir = directory / "english"
        _, _, english_manifest, english_path = artifact_set(english_dir, language="en")
        run("validate-test-points", english_dir / "test-points.json")
        run("render-test-points-xmind", english_dir / "test-points.json")
        run("render-test-points-xmind", english_dir / "test-points.json", "--format", "zen", "--output", english_dir / "points-zen.xmind")
        run("render-case-workbook", english_path)
        run("validate-case-workbook", english_path)
        run("import-case-workbook", english_path)
        imported = json.loads((english_dir / "test-manifest.pending.json").read_text(encoding="utf-8"))
        if imported != english_manifest:
            raise RuntimeError("English workbook roundtrip changed business data")
        english_book = load_workbook(english_dir / "test-cases.full.xlsx")
        try:
            if english_book["Test Cases"]["A1"].value != "Case ID":
                raise RuntimeError("English case workbook labels were not rendered")
        finally:
            english_book.close()
        run("freeze-coverage-plan", planned_path)
        frozen = json.loads(planned_path.read_text(encoding="utf-8"))
        if frozen["coverage_plan"]["status"] != "frozen" or frozen["requirement_confirmation"] != planned["requirement_confirmation"] or frozen["review"]["status"] != "draft":
            raise RuntimeError("coverage freezing must preserve human confirmation and leave semantic review pending")
        run("freeze-coverage-plan", directory / "design-context.json", fail=True)
        broken = copy.deepcopy(original_context)
        broken["rules"][0]["design_models"][0]["items"][0]["outcomes"].pop()
        write_json(directory / "changed-plan.json", broken)
        run("validate-design-context", directory / "changed-plan.json", fail=True)
        run("validate-design-context", directory / "design-context.json")
        run("validate-test-points", directory / "test-points.json")
        run("design-review-digest", directory / "test-points.json")
        run("design-review-digest", path)
        run("render-test-points-xmind", directory / "test-points.json")
        run("render-test-points-xmind", directory / "test-points.json", "--format", "zen", "--output", directory / "points-zen.xmind")
        run("validate-test-manifest", path)
        run("render-case-workbook", path)
        run("validate-case-workbook", path)
        # A new export changes the workbook container hash; bind this run to that export.
        workbook_hash = sha256_file(directory / "test-cases.smoke.xlsx")
        profile["case_workbook_sha256"] = workbook_hash
        write_json(profile_path, profile)
        results["run"]["case_workbook_sha256"] = workbook_hash
        results["run"]["execution_profile_sha256"] = sha256_file(profile_path)
        write_json(results_path, results)
        run("prepare-execution-profile", path, "--suite", "smoke", "--target", "web-chrome", "--output", results_path.parent / "missing-profile.json")
        run("validate-execution-profile", profile_path, path)
        run("validate-execution-results", results_path, path)
        run("judge-execution-results", results_path, path)
        run("validate-execution-results", results_path.with_name("execution-results.judged.json"), path)
        run("render-execution-report", results_path, path)
        # Missing planned checks must fail independently of a successful interaction.
        removed = results["results"][0]["assertions"].pop()
        write_json(results_path, results)
        run("validate-execution-results", results_path, path, fail=True)
        results["results"][0]["assertions"].append(removed)
        results["results"][0].update(status="failed", failure_reason="Controlled synthetic mismatch")
        removed.update(status="failed", actual="Synthetic mismatching outcome")
        mutate_observation(results["results"][0], removed, results_path.parent, "Synthetic mismatching outcome")
        # A falsely declared pass must be rejected, then judging must recover the
        # failing verdict from the evidence without changing the expected value.
        removed["status"] = "passed"
        results["results"][0]["status"] = "passed"
        write_json(results_path, results)
        run("validate-execution-results", results_path, path, fail=True)
        run("judge-execution-results", results_path, path, "--output", results_path.with_name("fault-judged.json"))
        run("validate-execution-results", results_path.with_name("fault-judged.json"), path)
        removed["status"] = "failed"
        results["results"][0]["status"] = "failed"
        write_json(results_path, results)
        run("render-execution-report", results_path, path)
        report = load_workbook(results_path.parent / "test-execution-report.xlsx")
        if report["执行汇总"]["B3"].value != "FAILED":
            raise RuntimeError("controlled mismatch did not fail the report gate")
        report.close()
        workbook = load_workbook(directory / "test-cases.smoke.xlsx")
        workbook["测试用例"]["C2"] = "Tester edit"
        workbook.save(directory / "test-cases.smoke.xlsx")
        workbook.close()
        run("validate-case-workbook", path, fail=True)
        run("import-case-workbook", path)
        pending = json.loads((directory / "test-manifest.pending.json").read_text(encoding="utf-8"))
        if pending["review"]["status"] != "pending" or not any(c["title"] == "Tester edit" for c in pending["cases"]):
            raise RuntimeError("Excel edit was not retained for review")
        run("validate-test-manifest", directory / "test-manifest.pending.json", fail=True)
        source = directory / "source.docx"
        from docx import Document
        document = Document()
        document.add_paragraph("Synthetic source rule")
        document.save(source)
        run("extract-source", source, "--source-id", "SRC-001", "--output", directory / "source.json")
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Synthetic flow"
        presentation.save(directory / "source.pptx")
        run("extract-source", directory / "source.pptx", "--source-id", "SRC-002", "--output", directory / "slides.json")
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(directory / "source.pdf")
        run("extract-source", directory / "source.pdf", "--source-id", "SRC-003", "--output", directory / "pdf.json")
        from PIL import Image
        Image.new("RGB", (10, 10), "white").save(directory / "source.png")
        run("extract-source", directory / "source.png", "--source-id", "SRC-004", "--output", directory / "image.json")
        run("extract-source", directory / "test-cases.smoke.xlsx", "--source-id", "SRC-005", "--output", directory / "excel-source.json")
        run("extract-source", directory / "test-points.xmind", "--source-id", "SRC-006", "--output", directory / "xmind-source.json")
        context_path, points_path = directory / "design-context.json", directory / "test-points.json"
        context = json.loads(context_path.read_text(encoding="utf-8"))
        context["review"] = {"status": "draft"}
        write_json(context_path, context)
        points = json.loads(points_path.read_text(encoding="utf-8"))
        points["design_context_baseline"]["sha256"] = sha256_file(context_path)
        seal_synthetic_points(points)
        write_json(points_path, points)
        run("validate-test-points", points_path, "--draft")
        run("render-test-points-xmind", points_path, "--draft")
        run("render-test-points-xmind", points_path, fail=True)
        # A review draft must not bypass unfinished requirement discovery.
        context["open_questions"] = [{"id":"Q-001", "unit_id":"RU-001", "impact":"Changes scope", "description":"Synthetic unanswered scope", "material":True}]
        write_json(context_path, context)
        points["design_context_baseline"]["sha256"] = sha256_file(context_path)
        seal_synthetic_points(points)
        write_json(points_path, points)
        run("validate-requirement-readiness", context_path, fail=True)
        run("render-test-points-xmind", points_path, "--draft", fail=True)
    print("OK: bundled native execution/summary, coverage-plan freezing, extraction, XMind, Excel roundtrip, profile, assertion rejection and failure report smoke")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

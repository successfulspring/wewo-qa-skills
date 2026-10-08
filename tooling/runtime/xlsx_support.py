from __future__ import annotations

import io
import math
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from artifact_tools import ValidationFailure


def text_cell(cell: Any, value: Any) -> None:
    """Write source content literally; spreadsheet formulas are never case data."""
    text = "" if value is None else str(value)
    if len(text) > 32767:
        raise ValidationFailure([f"{cell.coordinate}: Excel cell exceeds 32767 characters; split the content"])
    cell.value = text
    cell.data_type = "s"


def add_table(
    workbook: Workbook, name: str, headers: list[str], rows: Iterable[list[Any]],
    widths: dict[str, int] | None = None, choices: dict[str, list[str]] | None = None,
) -> Any:
    sheet = workbook.create_sheet(name)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    for index, header in enumerate(headers, 1):
        cell = sheet.cell(1, index)
        text_cell(cell, header)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="234E70")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = (widths or {}).get(header, 24)
    sheet.row_dimensions[1].height = 32
    for row_index, values in enumerate(rows, 2):
        max_lines = 1
        for index, value in enumerate(values, 1):
            cell = sheet.cell(row_index, index)
            text_cell(cell, value)
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F6FA")
            width = (widths or {}).get(headers[index - 1], 24)
            max_lines = max(max_lines, sum(max(1, math.ceil(len(line) * 1.4 / width)) for line in str(value or "").split("\n")))
        sheet.row_dimensions[row_index].height = min(409, max(30, max_lines * 16 + 8))
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for header, options in (choices or {}).items():
        column = get_column_letter(headers.index(header) + 1)
        validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        validation.error = "请从下拉列表选择值"
        validation.errorTitle = "无效值"
        validation.showErrorMessage = True
        validation.errorStyle = "stop"
        sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}1048576")
    return sheet


def workbook_bytes(workbook: Workbook) -> bytes:
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def load_case_excel(path: Path) -> Any:
    try:
        return load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValidationFailure([f"cannot read Excel workbook {path.name}: {type(exc).__name__}: {exc}"]) from exc


def read_table(workbook: Any, name: str, headers: list[str]) -> list[dict[str, str]]:
    if name not in workbook.sheetnames:
        raise ValidationFailure([f"missing worksheet {name}"])
    sheet = workbook[name]
    # Some document exports declare A1 even when the stored sheet has many rows.
    sheet.reset_dimensions()
    rows = iter(sheet.iter_rows())
    first = next(rows, ())
    actual = [str(c.value or "") for c in first]
    if actual != headers:
        raise ValidationFailure([f"{name}: headers changed; preserve all template columns"])
    result = []
    for row in rows:
        if not any(c.value is not None for c in row):
            continue
        if len(row) > len(headers) and any(c.value is not None for c in row[len(headers):]):
            raise ValidationFailure([f"{name}: unrecognized extra columns"])
        values = []
        for cell in row[:len(headers)]:
            if cell.data_type == "f":
                raise ValidationFailure([f"{name}!{cell.coordinate}: formulas are not allowed in editable case fields"])
            values.append("" if cell.value is None else str(cell.value))
        values += [""] * (len(headers) - len(values))
        result.append(dict(zip(headers, values)))
    return result

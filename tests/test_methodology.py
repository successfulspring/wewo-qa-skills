"""Maintenance tests for concrete omissions and false verdicts, not prose wording."""
from __future__ import annotations

import copy
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(ROOT / "tooling/runtime"), str(ROOT / "skills/wewo-qa-case-designer/scripts"), str(ROOT / "skills/wewo-qa-case-executor/scripts")]

from openpyxl import load_workbook
from artifact_factory import artifact_set, execution_set, mutate_observation, write_json
from artifact_tools import ValidationFailure, load_json, sha256_file, validate_results_data, validate_test_points_data
from case_workbook import read_case_workbook, comparison_cells, comparison_value
from design_context import validate_design_context_data
from design_methods import derivation_errors
from judge_execution_results import judge
from observation_checks import check_errors, compare


def model_context(method, items, **fields):
    context = load_json(ROOT / "tests/fixtures/design-context.json")
    rule = context["rules"][0]
    dimension = {"boundary":"boundary", "state-transition":"state", "permission":"role", "cross-object":"cross-object"}.get(method, "positive")
    for coverage in rule["coverage"]:
        coverage.update(applicability="required" if coverage["dimension"] == dimension else "not-applicable", test_point_refs=["TP-X"] if coverage["dimension"] == dimension else [], reason="Outside this isolated method test.")
    concrete = []
    for index, details in enumerate(items):
        item = {"id":f"COV-SAMPLE-{index}", "dimensions":[dimension], "condition":"Confirmed concrete condition", "action":"Perform the stated event", "expected":"The source-defined visible outcome", "setup":"Select a record in the specified state", "disposition":"required", "test_point_refs":[f"TP-SAMPLE-{index}"]}
        item.update(details)
        concrete.append(item)
    rule["design_models"] = [{"id":"DM-SAMPLE", "method":method, "rationale":"Isolated synthetic rule used to test omission detection", "items":concrete, **fields}]
    return context


class DerivationTests(unittest.TestCase):
    def test_numeric_boundary_omission_and_wrong_resolution(self):
        context = model_context("boundary", [{"boundary_ref":"MAX", "value":v} for v in [9,10,11]], field="Maximum count", bounds=[{"id":"MAX","value":10,"resolution":1,"basis":"At most 10 whole records"}])
        self.assertEqual(derivation_errors(context, final=True), [])
        model = context["rules"][0]["design_models"][0]
        model["items"].pop()
        self.assertTrue(any("below/on/above" in e for e in derivation_errors(context, final=True)))
        model["items"].append({**model["items"][0], "id":"COV-WRONG", "value":10.5})
        self.assertTrue(any("below/on/above" in e for e in derivation_errors(context, final=True)))

    def test_decimal_resolution_does_not_use_binary_float_rounding(self):
        context = model_context("boundary", [{"boundary_ref":"MAX", "value":v} for v in ["0.29","0.30","0.31"]], field="Price", bounds=[{"id":"MAX","value":"0.30","resolution":"0.01","basis":"One cent precision"}])
        self.assertEqual(derivation_errors(context, final=True), [])

    def test_decision_table_reduced_rows_cover_all_combinations(self):
        context = model_context("decision-table", [{"assignment":{"prepaid":"no","stock":None}},{"assignment":{"prepaid":"yes","stock":"no"}},{"assignment":{"prepaid":"yes","stock":"yes"}}], factors={"prepaid":["no","yes"],"stock":["no","yes"]})
        self.assertEqual(derivation_errors(context, final=True), [])
        context["rules"][0]["design_models"][0]["items"].pop()
        self.assertTrue(any("omits condition combinations" in e for e in derivation_errors(context, final=True)))

    def test_overlapping_rows_cannot_mask_missing_combination(self):
        context = model_context("decision-table", [{"assignment":{"a":"no","b":None}},{"assignment":{"a":None,"b":"no"}}], factors={"a":["no","yes"],"b":["no","yes"]})
        self.assertTrue(any("overlapping" in e for e in derivation_errors(context, final=True)))

    def test_infeasible_row_needs_reason_and_no_leaf(self):
        context = model_context("decision-table", [{"assignment":{"a":"no"}},{"assignment":{"a":"yes"},"disposition":"excluded","test_point_refs":[]}], factors={"a":["no","yes"]})
        self.assertTrue(any("excluded item needs reason" in e for e in derivation_errors(context, final=True)))
        context["rules"][0]["design_models"][0]["items"][1]["reason"] = "Confirmed process prevents this combination."
        self.assertEqual(derivation_errors(context, final=True), [])

    def test_state_matrix_requires_forbidden_event_pairs(self):
        context = model_context("state-transition", [{"from":s,"event":e,"to":s} for s in ["draft","submitted"] for e in ["submit","delete"]], object_ref="OBJ-ORDER", states=["draft","submitted"], events=["submit","delete"])
        context["objects"] = [{"id":"OBJ-ORDER","states":["draft","submitted"]}]
        context["rules"][0]["object_refs"] = ["OBJ-ORDER"]
        self.assertEqual(derivation_errors(context, final=True), [])
        context["rules"][0]["design_models"][0]["items"].pop()
        self.assertTrue(any("matrix must cover" in e for e in derivation_errors(context, final=True)))

    def test_transition_rule_matrix_does_not_require_fake_single_transition(self):
        context = model_context("state-transition", [{"from":s,"event":e,"to":s} for s in ["draft","submitted"] for e in ["submit","delete"]], object_ref="OBJ-ORDER", states=["draft","submitted"], events=["submit","delete"])
        context["objects"] = [{"id":"OBJ-ORDER","title":"Order","description":"Synthetic order states","states":["draft","submitted"],"segment_refs":["SEG-001-001"]}]
        context["rules"][0].update(kind="transition",object_refs=["OBJ-ORDER"])
        validate_design_context_data(context)

    def test_role_matrix_checks_denied_pairs(self):
        context = model_context("permission", [{"role":r,"operation":a} for r in ["viewer","editor"] for a in ["read","change"]], roles=["viewer","editor"], operations=["read","change"])
        self.assertEqual(derivation_errors(context, final=True), [])
        context["rules"][0]["design_models"][0]["items"].pop(1)
        self.assertTrue(any("matrix must cover" in e for e in derivation_errors(context, final=True)))

    def test_cross_object_requires_grounded_before_after_effects(self):
        context = model_context("cross-object", [{"effects":[{"object_ref":o,"before":"Bound to order O1","after":"Still bound to order O1"} for o in ["OBJ-A","OBJ-B"]]}])
        context["objects"] = [{"id":"OBJ-A"},{"id":"OBJ-B"}]
        context["rules"][0]["object_refs"] = ["OBJ-A","OBJ-B"]
        self.assertEqual(derivation_errors(context, final=True), [])
        context["rules"][0]["design_models"][0]["items"][0]["effects"][1]["object_ref"] = "OBJ-UNKNOWN"
        self.assertTrue(any("unrelated object" in e for e in derivation_errors(context, final=True)))

    def test_unresolved_review_is_not_a_confirmed_design(self):
        context = load_json(ROOT / "tests/fixtures/design-context.json")
        context["review"]["findings"] = [{"id":"F-001","category":"coverage","scope_refs":[context["rules"][0]["id"]],"description":"The wrong-password preservation expectation is missing","status":"open"}]
        with self.assertRaisesRegex(ValidationFailure, "unresolved review finding"):
            validate_design_context_data(context)
        context["review"]["findings"][0].update(status="resolved",resolution="Confirmed and added preservation test point.")
        validate_design_context_data(context)

    def test_coverage_link_and_dimension_mismatch_detected(self):
        with tempfile.TemporaryDirectory() as folder:
            context, points, _, _ = artifact_set(Path(folder))
            points["tree"]["children"][0]["children"][0]["coverage_item_refs"] = ["COV-UNKNOWN"]
            with self.assertRaisesRegex(ValidationFailure, "coverage item|leaf linkage"):
                validate_test_points_data(points, context)
            context["rules"][0]["design_models"][0]["items"][0]["test_point_refs"] = ["TP-WRONG-LEAF"]
            with self.assertRaisesRegex(ValidationFailure, "dimension leaf list"):
                validate_test_points_data(points, context)


class ObservationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.manifest, self.path, self.profile, _, self.results, self.result_path = execution_set(self.directory)
        self.result = self.results["results"][0]
        self.assertion = self.result["assertions"][0]

    def validate(self):
        validate_results_data(self.results, self.manifest, self.path, self.result_path.parent, self.profile)

    def record_path(self):
        eid = self.assertion["observations"][0]["evidence_id"]
        evidence = next(e for e in self.result["evidence"] if e["id"] == eid)
        return self.result_path.parent / evidence["path"], evidence

    def test_declared_pass_cannot_override_raw_mismatch(self):
        mutate_observation(self.result, self.assertion, self.result_path.parent, "Unexpected product state")
        self.assertion["actual"] = "Unexpected product state"
        with self.assertRaisesRegex(ValidationFailure, "declared status disagrees"):
            self.validate()
        judged = judge(self.results, self.manifest, self.result_path.parent)
        self.assertEqual(judged["results"][0]["status"], "failed")
        self.assertEqual(self.results["results"][0]["status"], "passed")
        validate_results_data(judged, self.manifest, self.path, self.result_path.parent, self.profile)

    def test_evidence_tampering_and_reused_run_rejected(self):
        path, evidence = self.record_path()
        record = load_json(path)
        record["run_id"] = "previous-run"
        write_json(path, record)
        with self.assertRaisesRegex(ValidationFailure, "hash mismatch"):
            self.validate()
        evidence["sha256"] = sha256_file(path)
        with self.assertRaisesRegex(ValidationFailure, "different run/case"):
            self.validate()

    def test_narrative_actual_cannot_substitute_original_output(self):
        self.assertion["actual"] = "Something the agent says it saw"
        with self.assertRaisesRegex(ValidationFailure, "actual text differs"):
            self.validate()

    def test_missing_pointer_and_stale_timestamp_rejected(self):
        self.assertion["observations"][0]["pointer"] = "/raw/nonexistent"
        with self.assertRaisesRegex(ValidationFailure, "pointer does not resolve"):
            self.validate()
        self.assertion["observations"][0]["pointer"] = "/raw/value"
        path, evidence = self.record_path()
        record = load_json(path)
        record["captured_at"] = "2026-09-03T10:00:30+08:00"
        write_json(path, record)
        evidence["sha256"] = sha256_file(path)
        with self.assertRaisesRegex(ValidationFailure, "outside this run"):
            self.validate()

    def test_second_passing_attempt_does_not_hide_first_failure(self):
        self.result["attempts"] = 2
        for index, assertion in enumerate(self.result["assertions"]):
            old = assertion["observations"][0]
            evidence = next(e for e in self.result["evidence"] if e["id"] == old["evidence_id"])
            record = load_json(self.result_path.parent / evidence["path"])
            second = {**evidence, "id":evidence["id"] + "-2", "path":evidence["path"].replace('.json','-2.json')}
            record["attempt"] = 2
            write_json(self.result_path.parent / second["path"], record)
            second["sha256"] = sha256_file(self.result_path.parent / second["path"])
            self.result["evidence"].append(second)
            assertion["evidence_refs"].append(second["id"])
            assertion["observations"].append({**old, "attempt":2,"evidence_id":second["id"]})
        path, evidence = self.record_path()
        first = load_json(path)
        first["raw"]["value"] = "Wrong state on first attempt"
        write_json(path, first)
        evidence["sha256"] = sha256_file(path)
        with self.assertRaisesRegex(ValidationFailure, "declared status disagrees|flaky"):
            self.validate()
        judged = judge(self.results, self.manifest, self.result_path.parent)
        self.assertEqual(judged["results"][0]["status"], "flaky")
        validate_results_data(judged, self.manifest, self.path, self.result_path.parent, self.profile)

    def test_evidence_review_requires_per_attempt_reasons(self):
        planned = self.manifest["cases"][0]["assertions"][0]
        planned["check"] = {"kind":"evidence-review","timing":"After rejection completes","criteria":"Reject with a message explaining submitted records cannot be deleted"}
        with self.assertRaisesRegex(ValidationFailure, "reasoned judgment"):
            self.validate()
        self.assertion["judgments"] = [{"attempt":1,"status":"passed","reason":"Original output states the confirmed prohibition; record remains shown."}]
        self.validate()

    def test_comparators_detect_wrong_delta_and_duplicate_loss(self):
        self.assertTrue(compare({"kind":"number-delta","expected_value":"0.10"}, "0.30", "0.20"))
        self.assertFalse(compare({"kind":"number-delta","expected_value":1}, 12, 10))
        self.assertFalse(compare({"kind":"unordered-equals","expected_value":["A","A","B"]}, ["B","A"]))
        self.assertTrue(compare({"kind":"unordered-equals","expected_value":["A","A","B"]}, ["B","A","A"]))
        self.assertFalse(compare({"kind":"equals","expected_value":True}, 1))
        self.assertFalse(compare({"kind":"unchanged"}, {"binding":"other"}, {"binding":"expected"}))
        self.assertTrue(check_errors({"kind":"number-equals","expected_value":"NaN","timing":"Ready"}))
        self.assertTrue(compare({"kind":"number-delta","expected_value":1}, "1234567890123456789012345678901", "1234567890123456789012345678900"))

    def test_before_after_must_have_same_subject_and_field(self):
        planned = self.manifest["cases"][0]["assertions"][0]
        planned["check"] = {"kind":"unchanged","timing":"After rejected change"}
        path, evidence = self.record_path()
        record = load_json(path)
        record["raw"] = {"value":"same","other":"same"}
        write_json(path,record)
        evidence["sha256"] = sha256_file(path)
        second = {**evidence,"id":"UI-BEFORE","path":"evidence/before.json"}
        record["subject"] = "another-record"
        write_json(self.result_path.parent / second["path"], record)
        second["sha256"] = sha256_file(self.result_path.parent / second["path"])
        self.result["evidence"].append(second)
        self.assertion["evidence_refs"].append(second["id"])
        self.assertion["observations"].insert(0,{"attempt":1,"phase":"before","evidence_id":second["id"],"pointer":"/raw/value"})
        self.assertion["actual"] = json.dumps({"after":"same","before":"same"},sort_keys=True)
        with self.assertRaisesRegex(ValidationFailure,"different subjects"):
            self.validate()
        record["subject"] = "synthetic-account"
        write_json(self.result_path.parent / second["path"],record)
        second["sha256"] = sha256_file(self.result_path.parent / second["path"])
        self.validate()
        self.assertion["observations"][0]["pointer"] = "/raw/other"
        with self.assertRaisesRegex(ValidationFailure,"same field pointer"):
            self.validate()

    def test_generated_excel_derivation_changes_are_not_silently_discarded(self):
        path = self.directory / "test-cases.xlsx"
        workbook = load_workbook(path)
        workbook["设计依据"]["J2"] = "Tester found a different expected outcome"
        workbook.save(path)
        workbook.close()
        with self.assertRaisesRegex(ValidationFailure, "generated design rows changed"):
            read_case_workbook(path)

    def test_comparison_edit_is_preserved_and_requires_review(self):
        path = self.directory / "test-cases.xlsx"
        workbook = load_workbook(path)
        workbook["必检断言"]["L2"] = "New source-confirmed value"
        workbook.save(path)
        workbook.close()
        candidate, original, _ = read_case_workbook(path)
        self.assertEqual(candidate["review"]["status"], "pending")
        self.assertEqual(candidate["review"]["checks"], [])
        self.assertEqual(candidate["cases"][0]["assertions"][0]["check"]["expected_value"], "New source-confirmed value")
        self.assertNotEqual(candidate, original)

    def test_tester_values_retain_empty_text_null_and_structured_types(self):
        for value in ["", "01", 'Text with "quotes"', 0, 1.25, False, None, ["A","A"], {"order":"O1","enabled":True}]:
            with self.subTest(value=value):
                cells = comparison_cells({"expected_value":value})
                self.assertEqual(comparison_value(*cells), {"expected_value":value})
        self.assertEqual(comparison_value(*comparison_cells({})), {})
        with self.assertRaises(ValidationFailure):
            comparison_value("数值", "true")


if __name__ == "__main__":
    unittest.main()

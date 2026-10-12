from __future__ import annotations

import copy
import io
import json
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

import test_workflow as workflow
from artifact_factory import write_json, seal_synthetic_cases, seal_synthetic_profile
from case_workbook import render_case_workbooks
from artifact_tools import ValidationFailure, sha256_file, validate_execution_profile_file
from prepare_execution_profile import build_profile
from native_execution import execute
from execution_summary import aggregate, load_summary
from render_execution_report import automatic_gate, render_report
from judge_execution_results import judge


class ExecutorDiscoveryTests(unittest.TestCase):
    setUp = workflow.NativeRunnerTests.setUp
    run_fixture = workflow.NativeRunnerTests.run_fixture

    def test_draft_discovery_cannot_start_a_runner(self):
        draft = build_profile(self.manifest,self.path,self.profile_path,"smoke",["web-chrome"])
        write_json(self.profile_path,draft)
        self.plan["profile_sha256"] = sha256_file(self.profile_path)
        with patch("native_execution.subprocess.run") as runner:
            with self.assertRaisesRegex(ValidationFailure,"readiness is pending"):
                execute(self.plan,self.path,self.profile_path)
            runner.assert_not_called()

    def test_discovered_need_blocks_only_its_ready_pair_until_resolved(self):
        requirement = {"id":"RT-EXTRA-DATA","name":"Prepared account","kind":"test-data","scope":"case",
                       "description":"Observed missing isolated account","sensitive":False,"collection":"user-input"}
        # Reuse an existing schema-supported data kind rather than invent one.
        requirement["kind"] = next(r["kind"] for r in self.manifest["runtime_requirements"] if not r["sensitive"])
        self.profile["discovered_requirements"].append(requirement)
        self.profile["case_readiness"][0]["runtime_requirement_refs"].append(requirement["id"])
        self.profile["bindings"].append({"requirement_id":requirement["id"],"status":"missing"})
        write_json(self.profile_path,self.profile)
        with self.assertRaisesRegex(ValidationFailure,"unresolved execution prerequisites"):
            validate_execution_profile_file(self.profile_path,self.path)
        self.profile["case_readiness"][0].update(status="blocked",reason="Isolated account is unavailable")
        write_json(self.profile_path,self.profile)
        validate_execution_profile_file(self.profile_path,self.path)
        summary,view,_ = aggregate(self.path,self.profile_path)
        self.assertEqual(view["results"][0]["status"],"blocked")
        self.assertIn("Isolated account",view["results"][0]["blocker"])
        self.assertEqual(automatic_gate(view,self.manifest),"INCOMPLETE")

    def test_ready_requires_observation_and_verified_environment(self):
        self.profile["case_readiness"][0]["checks"][1].update(status="not-applicable",detail="No observation route")
        self.profile["environments"][0]["verified"] = False
        write_json(self.profile_path,self.profile)
        with self.assertRaisesRegex(ValidationFailure,"verified target environment"):
            validate_execution_profile_file(self.profile_path,self.path)

    def test_runtime_route_cannot_replace_reviewed_ui_evidence(self):
        self.plan["bindings"][0]["route"] = "api"
        with self.assertRaisesRegex(ValidationFailure,"evidence surface"):
            execute(self.plan,self.path,self.profile_path)

    def test_independent_workspace_runs_without_product_repository(self):
        self.assertFalse((self.repo/".git").exists())
        _,results = self.run_fixture()
        self.assertEqual(results["results"][0]["status"],"passed")
        summary,view,_ = aggregate(self.path,self.profile_path)
        self.assertEqual(automatic_gate(view,self.manifest),"PASSED")
        self.assertEqual(len(summary["sources"]),1)

    def test_cleanup_failure_survives_judging_and_report_gate(self):
        self.runner.write_text(self.runner.read_text(encoding="utf-8").replace('data.get("cleanup_status","passed")','"failed"'),encoding="utf-8")
        self.plan["bindings"][0]["asset_sha256"] = sha256_file(self.runner)
        _,results = self.run_fixture()
        result = results["results"][0]
        self.assertTrue(all(a["status"] == "passed" for a in result["assertions"]))
        self.assertEqual(result["status"],"blocked")
        self.assertEqual(judge(results,self.manifest,self.profile_path.parent)["results"][0]["status"],"blocked")
        _,view,_ = aggregate(self.path,self.profile_path)
        self.assertEqual(automatic_gate(view,self.manifest),"INCOMPLETE")

    def test_wrong_actual_build_cannot_produce_business_pass(self):
        self.runner.write_text(self.runner.read_text(encoding="utf-8").replace('build=env["build"]','build="wrong-build"'),encoding="utf-8")
        self.plan["bindings"][0]["asset_sha256"] = sha256_file(self.runner)
        _,results = self.run_fixture()
        self.assertEqual(results["results"][0]["status"],"blocked")
        self.assertTrue(all(a["status"] == "not-evaluated" for a in results["results"][0]["assertions"]))

    def test_equivalent_failure_then_pass_is_flaky_in_whole_run(self):
        events = copy.deepcopy(self.events)
        events[0]["value"] = "Observed controlled mismatch"
        self.run_fixture(events=events)
        self.run_fixture()
        summary,view,_ = aggregate(self.path,self.profile_path)
        self.assertEqual(view["results"][0]["status"],"flaky")
        self.assertEqual(len(view["history"]),2)
        self.assertEqual(automatic_gate(view,self.manifest),"INCOMPLETE")
        book = load_workbook(io.BytesIO(render_report(view,self.manifest,results_dir=self.profile_path.parent,report_dir=self.root)))
        self.assertEqual(book["执行历史"].max_row,3)
        book.close()

    def test_missing_invocation_results_cannot_be_hidden(self):
        path,_ = self.run_fixture()
        path.unlink()
        with self.assertRaisesRegex(ValidationFailure,"no result file"):
            aggregate(self.path,self.profile_path)

    def test_summary_becomes_stale_after_another_invocation(self):
        self.run_fixture()
        summary,_,_ = aggregate(self.path,self.profile_path)
        path = self.profile_path.with_name("execution-summary.json")
        write_json(path,summary)
        self.run_fixture()
        with self.assertRaisesRegex(ValidationFailure,"stale or modified"):
            load_summary(path,self.path)

    def test_profile_cannot_be_rewritten_after_execution(self):
        self.run_fixture()
        self.profile["environments"][0]["build"] = "fixture-v2"
        write_json(self.profile_path,self.profile)
        self.plan["profile_sha256"] = sha256_file(self.profile_path)
        with self.assertRaisesRegex(ValidationFailure,"profile changed after execution"):
            execute(self.plan,self.path,self.profile_path)

    def test_malformed_lifecycle_still_produces_an_incomplete_report(self):
        self.runner.write_text(self.runner.read_text(encoding="utf-8").replace('data.get("cleanup_status","passed")','"invalid-status"'),encoding="utf-8")
        self.plan["bindings"][0]["asset_sha256"] = sha256_file(self.runner)
        _,results = self.run_fixture()
        self.assertEqual(results["results"][0]["status"],"blocked")
        _,view,_ = aggregate(self.path,self.profile_path)
        self.assertEqual(automatic_gate(view,self.manifest),"INCOMPLETE")

    def test_unavailable_runner_is_reported_as_blocked(self):
        self.plan["argv"] = [str(self.repo / "absent-runner.exe")]
        _,results = self.run_fixture()
        self.assertEqual(results["results"][0]["status"],"blocked")
        _,view,_ = aggregate(self.path,self.profile_path)
        self.assertEqual(automatic_gate(view,self.manifest),"INCOMPLETE")

    def test_missing_case_evidence_does_not_break_incomplete_reporting(self):
        self.case["automation"][0]["required_evidence"].append("screenshot")
        seal_synthetic_cases(self.manifest)
        write_json(self.path,self.manifest)
        render_case_workbooks(self.path)
        self.profile.update(manifest_sha256=sha256_file(self.path),case_workbook_sha256=sha256_file(self.root / "test-cases.smoke.xlsx"))
        write_json(self.profile_path,self.profile)
        self.plan.update(manifest_sha256=sha256_file(self.path),profile_sha256=sha256_file(self.profile_path))
        _,results = self.run_fixture()
        self.assertEqual(results["results"][0]["status"],"blocked")
        self.assertEqual(judge(results,self.manifest,self.profile_path.parent)["results"][0]["status"],"blocked")
        aggregate(self.path,self.profile_path)

    def test_missing_data_does_not_prevent_an_independent_ready_case(self):
        self.profile = build_profile(self.manifest,self.path,self.profile_path,"regression",["web-chrome"])
        old_bindings = {b["requirement_id"]:b for b in json.loads(self.profile_path.read_text())["bindings"]}
        for index,binding in enumerate(self.profile["bindings"]):
            if binding["requirement_id"] in old_bindings:
                self.profile["bindings"][index] = old_bindings[binding["requirement_id"]]
        seal_synthetic_profile(self.profile,self.environment)
        self.profile["case_readiness"][1].update(status="blocked",reason="Isolated lockout reset data is unavailable")
        write_json(self.profile_path,self.profile)
        self.plan["profile_sha256"] = sha256_file(self.profile_path)
        self.run_fixture()
        _,view,_ = aggregate(self.path,self.profile_path)
        self.assertEqual([r["status"] for r in view["results"]],["passed","blocked"])
        self.assertIn("lockout",view["results"][1]["blocker"])

    def test_report_fixed_labels_follow_english_without_translating_actual_values(self):
        _,results = self.run_fixture()
        manifest = copy.deepcopy(self.manifest)
        manifest["project"]["output_language"] = "en"
        book = load_workbook(io.BytesIO(render_report(results,manifest)))
        self.assertIn("Execution Summary",book.sheetnames)
        self.assertIn("Preparation and Cleanup",book.sheetnames)
        self.assertEqual(book["Assertion Results"]["G2"].value,results["results"][0]["assertions"][0]["actual"])
        for sheet in book:
            for cell in sheet[1]:
                self.assertFalse(any("\u3400" <= c <= "\u9fff" for c in str(cell.value)))
        book.close()


if __name__ == "__main__":
    unittest.main()

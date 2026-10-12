"""Bilingual artifact round trips, literal preservation and legacy compatibility."""
import copy
import io
import json
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(ROOT / p) for p in ("tooling/runtime", "skills/wewo-qa-case-designer/scripts", "skills/wewo-qa-case-executor/scripts", "tests")]
from openpyxl import load_workbook
from artifact_factory import artifact_set, seal_synthetic_points, seal_synthetic_cases
from artifact_tools import ValidationFailure, validate_test_points_data, validate_manifest_data
from artifact_language import output_language
from case_workbook import read_case_workbook, comparison_cells, comparison_value, render_case_workbook
from design_context import requirement_digest
from render_test_points_xmind import render_xmind_bytes


class ArtifactLanguageTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def english(self):
        return artifact_set(self.root, language="en")

    def test_english_headers_options_help_and_all_suite_roundtrips(self):
        _, _, manifest, _ = self.english()
        for suite in ("smoke", "regression", "full"):
            with self.subTest(suite=suite):
                path = self.root / f"test-cases.{suite}.xlsx"
                candidate, original, _ = read_case_workbook(path)
                self.assertEqual(candidate, manifest)
                self.assertEqual(original, manifest)
                book = load_workbook(path)
                try:
                    self.assertIn("Required Assertions", book.sheetnames)
                    self.assertEqual(book["Test Cases"]["A1"].value, "Case ID")
                    for sheet in book:
                        if sheet.title == "_baseline":
                            continue
                        self.assertIsNone(re.search(r"[\u3400-\u9fff]", sheet.title))
                        for cell in sheet[1]:
                            self.assertIsNone(re.search(r"[\u3400-\u9fff]", str(cell.value)))
                        for validation in sheet.data_validations.dataValidation:
                            self.assertIsNone(re.search(r"[\u3400-\u9fff]", validation.formula1 + validation.error + validation.errorTitle))
                    for row in book["Instructions"].iter_rows(min_row=3, values_only=True):
                        self.assertIsNone(re.search(r"[\u3400-\u9fff]", " ".join(map(str, row))))
                    # Authored source/test values remain exact even under English labels.
                    self.assertEqual(book["Test Cases"]["C2"].value,
                        next(c["title"] for c in manifest["cases"] if c["id"] == book["Test Cases"]["A2"].value))
                finally:
                    book.close()

    def test_english_edits_import_with_canonical_risk_and_pending_review(self):
        _, _, manifest, _ = self.english()
        path = self.root / "test-cases.full.xlsx"
        book = load_workbook(path)
        cid = book["Test Cases"]["A2"].value
        book["Test Cases"]["C2"] = "Updated English case"
        book["Test Cases"]["E2"] = "Low"
        book.save(path)
        book.close()
        candidate, original, _ = read_case_workbook(path)
        case = next(c for c in candidate["cases"] if c["id"] == cid)
        self.assertEqual((case["title"], case["risk"]), ("Updated English case", "low"))
        self.assertEqual(candidate["review"]["status"], "pending")
        self.assertEqual(original, manifest)
        self.assertEqual(candidate["sources"], manifest["sources"])

    def test_typed_values_and_literal_labels_are_not_translated(self):
        for language in ("zh-CN", "en"):
            for value in ("是", "Yes", "数值", "001", True, False, None, 12, 0.25, ["是", "Yes", 1], {"状态":"是"}):
                with self.subTest(language=language, value=value):
                    cells = comparison_cells({"expected_value":value}, language)
                    self.assertEqual(comparison_value(*cells, language), {"expected_value":value})
            self.assertEqual(comparison_value(*comparison_cells({}, language), language), {})

    def test_default_chinese_and_existing_format_13_still_import(self):
        _, _, manifest, _ = artifact_set(self.root)
        self.assertEqual(output_language(manifest), "zh-CN")
        path = self.root / "test-cases.full.xlsx"
        book = load_workbook(path)
        self.assertIn("测试用例", book.sheetnames)
        # Recreate the prior format's metadata; its Chinese sheets and values are unchanged.
        metadata = book["_baseline"]
        for row in list(metadata.iter_rows()):
            if row[0].value == "format":
                row[1].value = "1.3"
        for row in list(metadata.iter_rows()):
            if row[0].value == "output_language":
                metadata.delete_rows(row[0].row)
                break
        book.save(path)
        book.close()
        candidate, original, _ = read_case_workbook(path)
        self.assertEqual(candidate, manifest)
        self.assertEqual(original, manifest)

    def test_tampered_locale_cannot_reinterpret_an_english_workbook(self):
        self.english()
        path = self.root / "test-cases.full.xlsx"
        book = load_workbook(path)
        for row in book["_baseline"]:
            if row[0].value == "output_language":
                row[1].value = "zh-CN"
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValidationFailure, "language does not match"):
            read_case_workbook(path)

    def test_new_workbooks_require_recorded_language(self):
        self.english()
        path = self.root / "test-cases.full.xlsx"
        book = load_workbook(path)
        for row in list(book["_baseline"]):
            if row[0].value == "output_language":
                book["_baseline"].delete_rows(row[0].row)
                break
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValidationFailure, "unsupported project.output_language"):
            read_case_workbook(path)

    def test_english_formula_injection_remains_rejected(self):
        self.english()
        path = self.root / "test-cases.full.xlsx"
        book = load_workbook(path)
        book["Test Cases"]["C2"] = '=HYPERLINK("https://example.invalid", "case")'
        book.save(path)
        book.close()
        with self.assertRaisesRegex(ValidationFailure, "formulas are not allowed"):
            read_case_workbook(path)

    def test_language_must_agree_across_baselines(self):
        context, points, manifest, _ = self.english()
        points["project"]["output_language"] = "zh-CN"
        seal_synthetic_points(points)
        with self.assertRaisesRegex(ValidationFailure, "output_language"):
            validate_test_points_data(points, context)
        with self.assertRaisesRegex(ValidationFailure, "output_language"):
            validate_manifest_data(manifest, points)
        context["project"]["output_language"] = "zh-CN"
        with self.assertRaises(ValidationFailure):
            render_case_workbook(manifest, "0" * 64, {}, context)

    def test_language_change_invalidates_the_recorded_requirement_snapshot(self):
        context, _, _, _ = artifact_set(self.root)
        original = requirement_digest(context)
        context["project"]["output_language"] = "en"
        self.assertNotEqual(requirement_digest(context), original)

    def test_both_xmind_formats_localize_generated_text_and_preserve_literals(self):
        _, points, _, _ = self.english()
        points["project"]["requirement"] = "Synthetic feature"
        points["tree"]["title"] = "Synthetic test points"
        points["tree"]["notes"] = "Literal source value: 是"
        before = copy.deepcopy(points)
        for kind, content in (("legacy", "content.xml"), ("zen", "content.json")):
            for draft in (False, True):
                with self.subTest(kind=kind, draft=draft):
                    with zipfile.ZipFile(io.BytesIO(render_xmind_bytes(points, kind, draft=draft))) as archive:
                        text = archive.read(content).decode("utf-8")
                    self.assertIn("Synthetic feature Test Points", text)
                    self.assertIn("ID: TP-ROOT", text)
                    self.assertIn("Expected: ", text)
                    self.assertNotIn("编号：", text)
                    self.assertNotIn("自动化候选：", text)
                    self.assertIn("Literal source value: 是", text)
                    self.assertEqual("Draft (awaiting confirmation)" in text, draft)
        self.assertEqual(points, before)

    def test_unsupported_language_is_not_silently_exported_as_chinese(self):
        with self.assertRaisesRegex(ValidationFailure, "unsupported"):
            output_language({"project":{"output_language":"unsupported"}})

    def test_empty_english_smoke_view_remains_a_named_workbook(self):
        context, _, manifest, _ = self.english()
        for case in manifest["cases"]:
            case["suite_membership"] = ["full"]
        seal_synthetic_cases(manifest)
        book = load_workbook(io.BytesIO(render_case_workbook(manifest, "0" * 64, {}, context, "smoke")))
        try:
            self.assertEqual(book["Test Cases"].max_row, 1)
            self.assertIn("Smoke case file", book["Instructions"]["B3"].value)
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()

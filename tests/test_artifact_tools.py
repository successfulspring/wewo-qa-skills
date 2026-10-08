from __future__ import annotations

import copy
import sys
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(ROOT / "tooling/runtime"), str(ROOT / "skills/wewo-qa-case-designer/scripts"), str(ROOT / "skills/wewo-qa-case-executor/scripts")]

from openpyxl import load_workbook
from artifact_tools import ValidationFailure, sha256_file, validate_manifest_data, validate_manifest_file, validate_results_data, validate_results_file, validate_execution_profile_data, validate_execution_profile_file, validate_test_points_data
from artifact_factory import artifact_set, execution_set, write_json, mutate_observation
from case_workbook import read_case_workbook, render_case_workbook, validate_case_workbook_file
from design_context import validate_design_context_data
from render_execution_report import automatic_gate, render_report
from render_test_points_xmind import render_xmind_bytes


class ArtifactTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.context, self.points, self.manifest, self.path = artifact_set(self.directory)

    def test_grounded_chain_and_excel_roundtrip(self):
        validate_manifest_file(self.path)
        validate_case_workbook_file(self.directory / "test-cases.xlsx", self.path)
        candidate, original, digest = read_case_workbook(self.directory / "test-cases.xlsx")
        self.assertEqual(candidate, self.manifest)
        self.assertEqual(original, self.manifest)
        self.assertEqual(digest, sha256_file(self.path))

    def test_context_unread_requirement_rejected(self):
        self.context["segments"][0]["read_status"] = "needs-review"
        with self.assertRaisesRegex(ValidationFailure, "has not been read"):
            validate_design_context_data(self.context)

    def test_context_rule_gap_rejected(self):
        self.context["rules"] = self.context["rules"][1:]
        with self.assertRaisesRegex(ValidationFailure, "has no modeled rule"):
            validate_design_context_data(self.context)

    def test_context_required_boundary_cannot_be_empty(self):
        self.context["rules"][0]["coverage"][2].update(applicability="required", test_point_refs=[])
        with self.assertRaisesRegex(ValidationFailure, "required coverage has no test points"):
            validate_design_context_data(self.context)

    def test_context_unknown_relationship_rule_rejected(self):
        self.context["relationships"] = [{"id": "REL-UNKNOWN", "from_object": "OBJ-A", "to_object": "OBJ-B", "description": "Synthetic relation", "rule_refs": ["RULE-MISSING"]}]
        with self.assertRaisesRegex(ValidationFailure, "unknown object|unknown or excluded rule"):
            validate_design_context_data(self.context)

    def test_material_question_prevents_final_context(self):
        self.context["open_questions"] = [{"id": "Q-001", "description": "Unresolved business rule", "material": True}]
        with self.assertRaisesRegex(ValidationFailure, "unresolved material"):
            validate_design_context_data(self.context)
        validate_design_context_data(self.context, final=False)

    def test_draft_xmind_is_reviewable_but_not_executable(self):
        self.context["review"] = {"status": "draft"}
        self.context["open_questions"] = [{"id": "Q-001", "description": "Pending scope", "material": True}]
        write_json(self.directory / "design-context.json", self.context)
        self.points["design_context_baseline"]["sha256"] = sha256_file(self.directory / "design-context.json")
        write_json(self.directory / "test-points.json", self.points)
        from artifact_tools import validate_test_points_file
        validate_test_points_file(self.directory / "test-points.json", final=False)
        with zipfile.ZipFile(BytesIO(render_xmind_bytes(self.points, draft=True))) as archive:
            self.assertIn("草案（待确认）", archive.read("content.xml").decode())
        self.manifest["test_points_baseline"]["sha256"] = sha256_file(self.directory / "test-points.json")
        write_json(self.path, self.manifest)
        with self.assertRaisesRegex(ValidationFailure, "confirmed review|unresolved material"):
            validate_manifest_file(self.path)

    def test_point_requires_concrete_verification(self):
        del self.points["tree"]["children"][0]["children"][0]["verification"]
        with self.assertRaisesRegex(ValidationFailure, "verification"):
            validate_test_points_data(self.points)

    def test_point_unknown_source_anchor_rejected(self):
        self.points["tree"]["children"][0]["children"][0]["trace_refs"] = ["SRC-001#not-read"]
        with self.assertRaisesRegex(ValidationFailure, "not inventoried"):
            validate_test_points_data(self.points, self.context)

    def test_point_rule_coverage_must_be_bidirectional(self):
        self.points["tree"]["children"][0]["children"][0]["rule_refs"] = ["RULE-LOGIN-002"]
        with self.assertRaisesRegex(ValidationFailure, "coverage rule|missing from business"):
            validate_test_points_data(self.points, self.context)

    def test_context_digest_change_blocks_manifest(self):
        self.context["review"]["notes"] = "Updated understanding"
        write_json(self.directory / "design-context.json", self.context)
        with self.assertRaisesRegex(ValidationFailure, "sha256"):
            validate_manifest_file(self.path)

    def test_deep_tree_preserved_in_both_xmind_formats(self):
        original = copy.deepcopy(self.points["tree"]["children"][0]["children"][0])
        branch = original
        for i in range(22):
            branch = {"id": f"TP-GROUP-{i}", "title": f"Business group {i}", "kind": "feature", "status": "confirmed", "requirement_unit_refs": ["RU-001"], "trace_refs": ["SRC-001"], "coverage_item_refs": [], "children": [branch]}
        self.points["tree"]["children"][0]["children"][0] = branch
        validate_test_points_data(self.points, self.context)
        for fmt in ("legacy", "zen"):
            with zipfile.ZipFile(BytesIO(render_xmind_bytes(self.points, fmt))) as archive:
                data = archive.read("content.xml" if fmt == "legacy" else "content.json").decode()
                self.assertIn("Business group 21", data)
                self.assertIn(original["id"], data)
                self.assertIn(original["verification"]["expected"], data)

    def test_leaf_without_case_rejected(self):
        self.manifest["cases"].pop()
        with self.assertRaisesRegex(ValidationFailure, "leaf test point is not covered"):
            validate_manifest_data(self.manifest, self.points)

    def test_grouping_node_cannot_be_case_coverage(self):
        self.manifest["cases"][0]["test_point_refs"] = ["TP-LOGIN"]
        with self.assertRaisesRegex(ValidationFailure, "unknown or non-leaf"):
            validate_manifest_data(self.manifest, self.points)

    def test_case_source_anchor_must_be_read(self):
        self.manifest["cases"][0]["source_refs"] = ["SRC-001#unread-invented-section"]
        write_json(self.path, self.manifest)
        with self.assertRaisesRegex(ValidationFailure, "case source anchor"):
            validate_manifest_file(self.path)

    def test_assertion_must_bind_a_step(self):
        self.manifest["cases"][0]["steps"][0]["assertion_refs"] = []
        with self.assertRaisesRegex(ValidationFailure, "not assigned to a step"):
            validate_manifest_data(self.manifest)

    def test_each_target_needs_leaf_assertions(self):
        for a in self.manifest["cases"][0]["assertions"]:
            a["target_ids"] = ["web-chrome"]
        with self.assertRaisesRegex(ValidationFailure, "each target requires planned assertions"):
            validate_manifest_data(self.manifest)

    def test_suite_nesting(self):
        self.manifest["cases"][0]["suite_membership"] = ["smoke", "full"]
        with self.assertRaisesRegex(ValidationFailure, "smoke membership requires regression"):
            validate_manifest_data(self.manifest)

    def test_case_review_required(self):
        self.manifest["review"] = {"status": "pending"}
        with self.assertRaisesRegex(ValidationFailure, "review is pending"):
            validate_manifest_data(self.manifest)
        validate_manifest_data(self.manifest, final=False)

    def test_runtime_inputs_cannot_collect_plain_credentials(self):
        req = next(r for r in self.manifest["runtime_requirements"] if r["sensitive"])
        req["collection"] = "user-input"
        with self.assertRaisesRegex(ValidationFailure, "sensitive requirements"):
            validate_manifest_data(self.manifest)

    def test_unknown_runtime_requirement_rejected(self):
        self.manifest["cases"][0]["runtime_requirement_refs"].append("RT-UNKNOWN")
        with self.assertRaisesRegex(ValidationFailure, "unknown runtime requirement"):
            validate_manifest_data(self.manifest)


class ExcelTests(unittest.TestCase):
    setUp = ArtifactTests.setUp

    def edit(self, sheet, cell, value):
        workbook = load_workbook(self.directory / "test-cases.xlsx")
        workbook[sheet][cell] = value
        workbook.save(self.directory / "test-cases.xlsx")
        workbook.close()

    def test_excel_edit_imports_pending_and_blocks_execution(self):
        self.edit("测试用例", "C2", "Edited business case")
        candidate, original, _ = read_case_workbook(self.directory / "test-cases.xlsx")
        self.assertEqual(candidate["cases"][0]["title"], "Edited business case")
        self.assertEqual(candidate["review"]["status"], "pending")
        self.assertEqual(original, self.manifest)
        with self.assertRaisesRegex(ValidationFailure, "unreviewed edits"):
            validate_case_workbook_file(self.directory / "test-cases.xlsx", self.path)

    def test_formula_rejected_without_evaluation(self):
        self.edit("测试用例", "C2", '=HYPERLINK("https://example.invalid","case")')
        with self.assertRaisesRegex(ValidationFailure, "formula"):
            read_case_workbook(self.directory / "test-cases.xlsx")

    def test_stale_workbook_blocked(self):
        self.manifest["review"]["notes"] = "A new baseline"
        write_json(self.path, self.manifest)
        with self.assertRaisesRegex(ValidationFailure, "stale or different"):
            validate_case_workbook_file(self.directory / "test-cases.xlsx", self.path)

    def test_target_feasibility_edit_retained(self):
        self.edit("自动化评估", "C2", "有条件自动化")
        self.edit("自动化评估", "F2", "Authorized route must be provided")
        candidate, _, _ = read_case_workbook(self.directory / "test-cases.xlsx")
        self.assertEqual(candidate["cases"][0]["automation"][0]["feasibility"], "conditional")
        self.assertEqual(candidate["review"]["status"], "pending")

    def test_multiline_data_and_large_path_metadata_lossless(self):
        self.manifest["cases"][0]["test_data"] = ["row one\nrow two\n", "another item  "]
        paths = {f"TP-{i}": "Long business path / " * 100 for i in range(100)}
        (self.directory / "test-cases.xlsx").write_bytes(render_case_workbook(self.manifest, sha256_file(self.path), paths, self.context))
        candidate, _, _ = read_case_workbook(self.directory / "test-cases.xlsx")
        self.assertEqual(candidate, self.manifest)

    def test_sorted_rows_are_not_business_edits(self):
        workbook = load_workbook(self.directory / "test-cases.xlsx")
        sheet = workbook["测试用例"]
        rows = [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]
        for i, row in enumerate(reversed(rows), 2):
            for j, value in enumerate(row, 1):
                sheet.cell(i, j).value = value
        workbook.save(self.directory / "test-cases.xlsx")
        workbook.close()
        candidate, _, _ = read_case_workbook(self.directory / "test-cases.xlsx")
        self.assertEqual(candidate, self.manifest)


class ExecutionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)
        self.manifest, self.path, self.profile, self.profile_path, self.results, self.results_path = execution_set(self.directory)

    def validate(self):
        validate_results_data(self.results, self.manifest, self.path, self.results_path.parent, self.profile)

    def test_complete_required_assertion_set_passes(self):
        self.validate()
        validate_results_file(self.results_path, self.path)
        self.assertEqual(automatic_gate(self.results, self.manifest), "PASSED")

    def test_omitted_required_assertion_cannot_pass(self):
        self.results["results"][0]["assertions"].pop()
        with self.assertRaisesRegex(ValidationFailure, "missing required assertion"):
            self.validate()

    def test_rewritten_oracle_rejected(self):
        self.results["results"][0]["assertions"][0]["expected"] = "whatever appeared"
        with self.assertRaisesRegex(ValidationFailure, "differs from planned oracle"):
            self.validate()

    def test_assertion_evidence_type_must_match(self):
        self.results["results"][0]["assertions"][0]["evidence_refs"] = [self.results["results"][0]["evidence"][0]["id"]]
        with self.assertRaisesRegex(ValidationFailure, "missing linked required evidence"):
            self.validate()

    def test_unknown_assertion_and_duplicate_evidence_rejected(self):
        result = self.results["results"][0]
        result["assertions"][0]["assertion_id"] = "AS-UNPLANNED"
        result["evidence"].append(copy.deepcopy(result["evidence"][0]))
        with self.assertRaisesRegex(ValidationFailure, "unknown or inapplicable|duplicate evidence"):
            self.validate()

    def test_unchecked_assertion_cannot_pass(self):
        self.results["results"][0]["assertions"][0].update(status="not-evaluated", actual="", evidence_refs=[], reason="Interrupted")
        with self.assertRaisesRegex(ValidationFailure, "requires only passing assertions"):
            self.validate()

    def test_product_mismatch_report_is_failed(self):
        result = self.results["results"][0]
        result.update(status="failed", failure_reason="Home nickname differs")
        result["assertions"][1].update(status="failed", actual="Another member nickname")
        mutate_observation(result, result["assertions"][1], self.results_path.parent, "Another member nickname")
        self.validate()
        workbook = load_workbook(BytesIO(render_report(self.results, self.manifest)))
        summary = {r[0].value: r[1].value for r in workbook["执行汇总"].iter_rows(min_row=2)}
        self.assertEqual(summary["执行门禁"], "FAILED")
        self.assertEqual(summary["断言检查覆盖率"], "2/2")
        self.assertEqual(workbook["断言结果"]["H3"].value, "失败")
        workbook.close()

    def test_blocked_zero_attempt_reports_all_unchecked_assertions(self):
        self.results["results"][0].update(status="blocked", attempts=0, assertions=[], evidence=[], blocker="Device unavailable")
        self.validate()
        workbook = load_workbook(BytesIO(render_report(self.results, self.manifest)))
        self.assertEqual(workbook["断言结果"].max_row, 3)
        self.assertEqual(workbook["断言结果"]["H2"].value, "未检查")
        self.assertEqual(automatic_gate(self.results, self.manifest), "INCOMPLETE")
        workbook.close()

    def test_workbook_change_invalidates_frozen_profile(self):
        with (self.directory / "test-cases.xlsx").open("ab") as stream:
            stream.write(b"changed")
        with self.assertRaisesRegex(ValidationFailure, "case_workbook_sha256"):
            validate_execution_profile_file(self.profile_path, self.path)

    def test_missing_evidence_file_rejected(self):
        (self.results_path.parent / self.results["results"][0]["evidence"][0]["path"]).unlink()
        with self.assertRaisesRegex(ValidationFailure, "evidence file does not exist"):
            self.validate()

    def test_evidence_escape_rejected(self):
        self.results["results"][0]["evidence"][0]["path"] = "../outside.txt"
        with self.assertRaisesRegex(ValidationFailure, "escapes run directory"):
            self.validate()

    def test_missing_runtime_condition_blocks_verdict(self):
        binding = self.profile["bindings"][0]
        self.profile["bindings"][0] = {"requirement_id": binding["requirement_id"], "status": "missing"}
        with self.assertRaisesRegex(ValidationFailure, "unresolved runtime requirement"):
            self.validate()

    def test_profile_rejects_plaintext_sensitive_value(self):
        binding = next(b for b in self.profile["bindings"] if "session_ref" in b)
        binding["value"] = "synthetic-value-that-must-not-be-stored"
        with self.assertRaisesRegex(ValidationFailure, "sensitive value must not be stored"):
            validate_execution_profile_data(self.profile, self.manifest, self.path)

    def test_profile_aggregates_selected_automatic_scope(self):
        from artifact_tools import expected_runtime_requirement_ids
        selected = expected_runtime_requirement_ids(self.manifest, "regression", ["ios-app"])
        self.assertNotIn("RT-DATA-LOCKOUT-RESET", selected)
        self.assertIn("RT-DEVICE-IOS", selected)
        self.assertEqual(len(selected), 5)

    def test_no_selected_pairs_cannot_pass_gate(self):
        self.results["results"] = []
        with self.assertRaisesRegex(ValidationFailure, "missing result|non-empty"):
            self.validate()
        self.assertEqual(automatic_gate(self.results, self.manifest), "NO_AUTOMATABLE_CASES")


if __name__ == "__main__":
    unittest.main()
